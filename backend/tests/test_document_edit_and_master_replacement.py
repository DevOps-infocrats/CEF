import io
import os
import uuid
import docx
from decimal import Decimal
from app.extensions import db
from app.modules.auth.models import (
    User,
    Role,
    UserRole,
    UserPermission,
    Permission,
)
from app.modules.auth.services import hash_password
from app.modules.candidates.models import (
    CandidateApplication,
    APPLICATION_STATUS_APPROVED,
)
from app.modules.letters.models import (
    DocumentTemplate,
    CandidateLetter,
    TEMPLATE_TYPE_LOI,
    TEMPLATE_TYPE_OFFER,
)
from app.modules.letters.services import (
    LetterTemplateService,
    LetterGenerationService,
    resolve_storage_path,
)


def create_user(email: str, password: str, permissions: list[str] = None, role: str = "staff"):
    user = User(
        email=email.lower().strip(),
        username=email.split("@")[0],
        password_hash=hash_password(password),
        first_name="Test",
        last_name="User",
        is_active=True,
        is_verified=True,
        approval_status="APPROVED",
    )
    db.session.add(user)
    db.session.flush()

    r = Role.query.filter_by(name=role).first()
    if r:
        ur = UserRole(user_id=user.id, role_id=r.id)
        db.session.add(ur)

    if permissions:
        for p_code in permissions:
            perm = Permission.query.filter_by(code=p_code).first()
            if perm:
                up = UserPermission(user_id=user.id, permission_id=perm.id)
                db.session.add(up)

    db.session.commit()
    return user


def login(client, email: str, password: str):
    resp = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert resp.status_code == 200, f"Login failed for {email}: {resp.json}"
    return {"Authorization": f"Bearer {resp.json['data']['access_token']}"}


def create_approved_candidate(user: User, full_name: str, app_num: str) -> CandidateApplication:
    app = CandidateApplication(
        application_number=app_num,
        user_id=user.id,
        hero_category="HDD",
        role_applied="HDD Operator",
        full_name=full_name,
        mobile_number="9876543210",
        email=f"{app_num.lower()}@example.com",
        aadhaar_number="987654321012",
        pan_number="ABCDE1234F",
        current_address_line1="123 Civil Lines",
        current_city="Jaipur",
        current_state="Rajasthan",
        current_pincode="302001",
        permanent_address_line1="123 Civil Lines",
        permanent_city="Jaipur",
        permanent_state="Rajasthan",
        permanent_pincode="302001",
        account_holder_name=full_name,
        bank_name="State Bank of India",
        account_number="12345678901",
        ifsc_code="SBIN0001234",
        circle="Rajasthan",
        district="Jaipur",
        block="Amber",
        current_salary=Decimal("35000.00"),
        asking_salary=Decimal("42000.00"),
        negotiated_salary=Decimal("40000.00"),
        status=APPLICATION_STATUS_APPROVED,
    )
    db.session.add(app)
    db.session.commit()
    return app


def test_document_editing_persistence_and_download(client):
    """
    Test Issue 1:
    - Open candidate document
    - Edit multiple fields (joining date, salary, intro text, custom paragraphs)
    - Save edits and verify HTTP 200 OK
    - Refresh/reopen document via GET API and confirm edits are preserved
    - Download document and verify that the actual downloaded DOCX contains the latest saved edits
    - Verify download headers prevent caching (no-cache, max-age=0)
    """
    approver = create_user("editor_hr@example.com", "Pass@1234", ["candidates:approve", "documents:view", "letters:generate"])
    headers = login(client, "editor_hr@example.com", "Pass@1234")

    cand_user = create_user("cand_edit_test@example.com", "Pass@1234", ["portal:access"], role="candidate")
    app = create_approved_candidate(cand_user, "Aarav Sharma", f"APP-EDIT-{uuid.uuid4().hex[:6].upper()}")

    # 1. Initialize candidate letters
    init_resp = client.get(f"/api/v1/applications/{app.id}/letters", headers=headers)
    assert init_resp.status_code == 200
    letters = init_resp.json["data"]
    assert "letter_of_intent" in letters
    assert "offer_letter" in letters

    # 2. Edit multiple fields & semantic text paragraphs on LOI
    custom_intro = "We are exceptionally thrilled to engage your services as Senior HDD Lead Specialist."
    custom_perf_note = "*Custom Performance Bonus: 15% guaranteed upon completion of quarter 1 targets."
    save_loi_resp = client.patch(
        f"/api/v1/applications/{app.id}/letters/letter_of_intent",
        headers=headers,
        json={
            "field_values": {
                "joining_date": "25/12/2026",
                "monthly_salary": "48000",
                "content_edits": {
                    "p_intro": custom_intro,
                    "p_perf_note": custom_perf_note,
                },
            },
            "text_edits": {
                "p_intro": custom_intro,
                "p_perf_note": custom_perf_note,
            },
        },
    )
    assert save_loi_resp.status_code == 200, f"Save failed: {save_loi_resp.json}"
    saved_loi = save_loi_resp.json["data"]
    assert saved_loi["field_values"]["joining_date"] == "25/12/2026"
    assert saved_loi["field_values"]["monthly_salary"] == "48000"
    assert saved_loi["field_values"]["content_edits"]["p_intro"] == custom_intro

    # 3. Refresh / reopen the document via GET API (simulating user page refresh)
    reload_resp = client.get(f"/api/v1/applications/{app.id}/letters", headers=headers)
    assert reload_resp.status_code == 200
    reloaded_letters = reload_resp.json["data"]
    reloaded_loi = reloaded_letters["letter_of_intent"]

    # Confirm edits remain visible after refresh
    assert reloaded_loi["field_values"]["joining_date"] == "25/12/2026"
    assert reloaded_loi["field_values"]["monthly_salary"] == "48000"
    assert reloaded_loi["field_values"]["content_edits"]["p_intro"] == custom_intro
    assert reloaded_loi["content"]["content_edits"]["p_intro"] == custom_intro

    # 4. Download candidate document and verify physical DOCX file content
    dl_resp = client.get(f"/api/v1/applications/{app.id}/letters/letter_of_intent/download?format=docx", headers=headers)
    assert dl_resp.status_code == 200
    assert "no-cache" in dl_resp.headers.get("Cache-Control", "")

    # Parse the downloaded DOCX file directly
    downloaded_doc = docx.Document(io.BytesIO(dl_resp.data))
    doc_text = "\n".join(p.text for p in downloaded_doc.paragraphs)

    assert custom_intro in doc_text, f"Custom intro not found in downloaded docx:\n{doc_text}"
    assert custom_perf_note in doc_text, f"Custom performance note not found in downloaded docx:\n{doc_text}"
    assert "25/12/2026" in doc_text, "Updated joining date not found in downloaded docx"


def test_master_template_replacement_and_subsequent_generation(client):
    """
    Test Issue 2:
    - Confirm active master template v1 exists
    - Upload a different valid master template (v2) with unique content and placeholders
    - Confirm the new master template is active (version >= 2, is_active=True)
    - Download master template and verify it downloads the uploaded v2 file
    - Generate a document for a NEW candidate and verify it uses the new master template
    - Confirm previously generated candidate documents remain unchanged
    - Test upload failure (corrupted file) and verify existing working master is preserved
    """
    admin = create_user("master_admin@example.com", "Pass@1234", ["users:manage", "candidates:approve"], role="admin")
    admin_headers = login(client, "master_admin@example.com", "Pass@1234")

    # 1. Existing candidate 1 generates document using initial master
    cand1 = create_user("cand_prev@example.com", "Pass@1234", ["portal:access"], role="candidate")
    app1 = create_approved_candidate(cand1, "Vikas Gupta", f"APP-PREV-{uuid.uuid4().hex[:6].upper()}")

    resp_c1 = client.get(f"/api/v1/applications/{app1.id}/letters", headers=admin_headers)
    assert resp_c1.status_code == 200
    loi1_initial = resp_c1.json["data"]["letter_of_intent"]

    # 2. Admin uploads a new master template for Letter of Intent (v2) with distinct text
    v2_doc = docx.Document()
    v2_doc.add_heading("FIBERFORCE MASTER LOI VERSION TWO", 0)
    v2_doc.add_paragraph("Welcome to the Next Generation Team, {{candidate_full_name}}!")
    v2_doc.add_paragraph("Official Posting Location: {{approved_location}}")
    v2_doc.add_paragraph("Start Date: {{joining_date}}")
    v2_doc.add_paragraph("Special Master v2 Clause: Highly confidential proprietary engagement terms.")

    buf = io.BytesIO()
    v2_doc.save(buf)
    buf.seek(0)

    upload_resp = client.post(
        "/api/v1/letters/templates/letter_of_intent/upload",
        headers=admin_headers,
        data={"file": (buf, "master_loi_revised_v2.docx")},
        content_type="multipart/form-data",
    )
    assert upload_resp.status_code == 201
    tmpl_data = upload_resp.json["data"]
    v2_version = tmpl_data["version"]
    assert v2_version >= 2
    assert tmpl_data["is_active"] is True

    # 3. Confirm active master template resolution returns v2
    active_meta = client.get("/api/v1/letters/templates/letter_of_intent", headers=admin_headers)
    assert active_meta.status_code == 200
    assert active_meta.json["data"]["version"] == v2_version
    assert active_meta.json["data"]["is_active"] is True

    # 4. Download master template and verify it is indeed v2
    dl_tmpl = client.get("/api/v1/letters/templates/letter_of_intent/download", headers=admin_headers)
    assert dl_tmpl.status_code == 200
    assert "no-cache" in dl_tmpl.headers.get("Cache-Control", "")
    downloaded_tmpl_doc = docx.Document(io.BytesIO(dl_tmpl.data))
    tmpl_text = "\n".join(p.text for p in downloaded_tmpl_doc.paragraphs)
    assert "FIBERFORCE MASTER LOI VERSION TWO" in tmpl_text
    assert "Special Master v2 Clause" in tmpl_text

    # 5. Subsequent document generation: Create Candidate 2 and generate their letters
    cand2 = create_user("cand_new_gen@example.com", "Pass@1234", ["portal:access"], role="candidate")
    app2 = create_approved_candidate(cand2, "Sneha Verma", f"APP-NEW-{uuid.uuid4().hex[:6].upper()}")

    resp_c2 = client.get(f"/api/v1/applications/{app2.id}/letters", headers=admin_headers)
    assert resp_c2.status_code == 200
    loi2 = resp_c2.json["data"]["letter_of_intent"]

    # Download Candidate 2's generated letter and verify it was created from master v2
    dl_c2 = client.get(f"/api/v1/applications/{app2.id}/letters/letter_of_intent/download?format=docx", headers=admin_headers)
    assert dl_c2.status_code == 200
    doc_c2 = docx.Document(io.BytesIO(dl_c2.data))
    c2_text = "\n".join(p.text for p in doc_c2.paragraphs)

    assert "FIBERFORCE MASTER LOI VERSION TWO" in c2_text, "Candidate 2 doc was not generated from Master v2"
    assert "Sneha Verma" in c2_text, "Candidate 2 legal name was not populated into Master v2"
    assert "Special Master v2 Clause" in c2_text, "Master v2 clause missing in Candidate 2 document"

    # 6. Verify Candidate 1's previously generated document remains unchanged
    dl_c1 = client.get(f"/api/v1/applications/{app1.id}/letters/letter_of_intent/download?format=docx", headers=admin_headers)
    assert dl_c1.status_code == 200
    doc_c1 = docx.Document(io.BytesIO(dl_c1.data))
    c1_text = "\n".join(p.text for p in doc_c1.paragraphs)
    assert "FIBERFORCE MASTER LOI VERSION TWO" not in c1_text, "Candidate 1 document was incorrectly overwritten by Master v2"

    # 7. Test upload failure (corrupted file)
    corrupted_resp = client.post(
        "/api/v1/letters/templates/letter_of_intent/upload",
        headers=admin_headers,
        data={"file": (io.BytesIO(b"corrupted binary data"), "bad.docx")},
        content_type="multipart/form-data",
    )
    assert corrupted_resp.status_code == 400

    # Confirm master v2 remains active and unaffected
    post_corrupt_meta = client.get("/api/v1/letters/templates/letter_of_intent", headers=admin_headers)
    assert post_corrupt_meta.status_code == 200
    assert post_corrupt_meta.json["data"]["version"] == v2_version
    assert post_corrupt_meta.json["data"]["is_active"] is True


import io
import uuid
from datetime import datetime, timezone
from decimal import Decimal
import openpyxl
import pytest
from app.extensions import db
from app.modules.auth.models import User, Role, UserRole
from app.modules.auth.services import hash_password
from app.modules.candidates.models import (
    CandidateApplication,
    APPLICATION_STATUS_APPROVED,
    APPLICATION_STATUS_PENDING,
    JOINING_STATUS_NOT_JOINED,
    JOINING_STATUS_JOINED,
)
from app.modules.letters.models import CandidateLetter, LETTER_STATUS_DRAFT, LETTER_STATUS_SENT


def create_user(email: str, password: str, first_name: str, last_name: str, role_name: str = "admin") -> User:
    """Helper to create and activate a test user with assigned role."""
    user = User(
        email=email,
        username=email.split("@")[0],
        password_hash=hash_password(password),
        first_name=first_name,
        last_name=last_name,
        is_active=True,
        is_verified=True,
        approval_status="APPROVED",
    )
    role = Role.query.filter_by(name=role_name).first()
    if not role:
        role = Role(name=role_name, description=f"{role_name} role")
        db.session.add(role)
        db.session.flush()

    db.session.add(user)
    db.session.flush()
    ur = UserRole(user_id=user.id, role_id=role.id)
    db.session.add(ur)
    db.session.commit()
    return user


def login(client, email: str, password: str) -> dict:
    """Helper to login and get Authorization bearer header."""
    res = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert res.status_code == 200, f"Login failed for {email}: {res.json}"
    token = res.json["data"]["access_token"]
    return {"Authorization": f"Bearer {token}"}


def test_approved_tab_and_editing_persistence(client, app):
    """Test Requirement 1 & 3: Approved candidate tracking, authorized editing, and preservation of audit/creator."""
    with app.app_context():
        hr_creator = create_user("creator_hr@test.com", "Pass@1234", "Pooja", "Nair", "reviewer")
        candidate_user = create_user("cand_appr@test.com", "Pass@1234", "Rohan", "Mehra", "candidate")

        app_rec = CandidateApplication(
            application_number="CEF-TEST-APP-001",
            user_id=candidate_user.id,
            created_by_id=hr_creator.id,
            hero_category="Hero 1",
            role_applied="HDD Operate",
            full_name="Rohan Mehra",
            mobile_number="9876543210",
            email="cand_appr@test.com",
            current_city="Jaipur",
            current_state="Rajasthan",
            current_salary=Decimal("25000.00"),
            status=APPLICATION_STATUS_APPROVED,
        )
        app_rec.reviewed_at = datetime.now(timezone.utc)
        app_rec.reviewed_by_id = hr_creator.id
        db.session.add(app_rec)
        db.session.commit()
        app_id = str(app_rec.id)

    headers_hr = login(client, "creator_hr@test.com", "Pass@1234")

    # 1. Verify candidate details retrieved via API
    get_res = client.get(f"/api/v1/applications/{app_id}", headers=headers_hr)
    assert get_res.status_code == 200
    data = get_res.json["data"]
    assert data["status"] == "APPROVED"
    assert data["created_by_name"] == "Pooja Nair"
    assert data["full_name"] == "Rohan Mehra"
    assert data["site_joining_status"] == "NOT_JOINED"
    assert data["enrollment_completed"] is False

    # 2. Authorized edit of candidate details
    patch_payload = {
        "full_name": "Rohan A. Mehra",
        "current_city": "Ajmer",
        "negotiated_salary": "28000.00",
        "joining_date": "20/11/2026",
    }
    patch_res = client.patch(f"/api/v1/applications/{app_id}", headers=headers_hr, json=patch_payload)
    assert patch_res.status_code == 200
    updated_data = patch_res.json["data"]
    assert updated_data["full_name"] == "Rohan A. Mehra"
    assert updated_data["current_city"] == "Ajmer"
    assert updated_data["negotiated_salary"] == "28000.00"
    assert updated_data["joining_date"] == "20/11/2026"
    # Creator must remain original Pooja Nair
    assert updated_data["created_by_name"] == "Pooja Nair"
    assert updated_data["application_number"] == "CEF-TEST-APP-001"

    # 3. Unauthorized candidate cannot edit approved record
    headers_cand = login(client, "cand_appr@test.com", "Pass@1234")
    unauth_patch = client.patch(
        f"/api/v1/applications/{app_id}",
        headers=headers_cand,
        json={"full_name": "Hacked Name"},
    )
    # The candidate user lacks candidates:write/approve
    assert unauth_patch.status_code in (401, 403)


def test_offer_letter_released_tracking_logic(client, app):
    """Test Requirement 1 & 7: Approved candidate does NOT appear as released until offer letter is actually sent."""
    with app.app_context():
        admin = create_user("admin_offer@test.com", "Pass@1234", "Arun", "Sharma", "admin")
        admin_id = admin.id
        cand = create_user("cand_offer@test.com", "Pass@1234", "Kavita", "Singh", "candidate")

        app_rec = CandidateApplication(
            application_number="CEF-TEST-OFFER-001",
            user_id=cand.id,
            created_by_id=admin_id,
            role_applied="HDD Tracker",
            full_name="Kavita Singh",
            email="cand_offer@test.com",
            status=APPLICATION_STATUS_APPROVED,
        )
        app_rec.reviewed_at = datetime.now(timezone.utc)
        db.session.add(app_rec)
        db.session.commit()
        app_id = str(app_rec.id)

    headers_admin = login(client, "admin_offer@test.com", "Pass@1234")

    # Step A: Approved candidate with no offer letter
    res1 = client.get(f"/api/v1/applications/{app_id}", headers=headers_admin)
    assert res1.status_code == 200
    assert res1.json["data"]["offer_letter_status"] == "NOT_GENERATED"
    assert res1.json["data"]["offer_letter_released_at"] is None

    # Step B: Offer letter in DRAFT status (generated but unsent)
    with app.app_context():
        draft_letter = CandidateLetter(
            application_id=uuid.UUID(app_id),
            letter_type="offer_letter",
            docx_storage_path="storage/letters/dummy.docx",
            status=LETTER_STATUS_DRAFT,
            sent_at=None,
        )
        db.session.add(draft_letter)
        db.session.commit()

    res2 = client.get(f"/api/v1/applications/{app_id}", headers=headers_admin)
    assert res2.status_code == 200
    assert res2.json["data"]["offer_letter_status"] == "DRAFT"
    assert res2.json["data"]["offer_letter_released_at"] is None

    # Step C: Offer letter actually SENT / RELEASED
    with app.app_context():
        letter = CandidateLetter.query.filter_by(application_id=uuid.UUID(app_id), letter_type="offer_letter").first()
        letter.status = LETTER_STATUS_SENT
        letter.sent_at = datetime.now(timezone.utc)
        letter.sent_by_id = admin_id
        db.session.commit()

    res3 = client.get(f"/api/v1/applications/{app_id}", headers=headers_admin)
    assert res3.status_code == 200
    assert res3.json["data"]["offer_letter_status"] == "RELEASED"
    assert res3.json["data"]["offer_letter_released_at"] is not None
    assert res3.json["data"]["offer_letter_released_by_name"] == "Arun Sharma"


def test_site_joining_manual_action_and_enrollment_completion(client, app):
    """Test Requirement 2 & 7: Manual site joining update, confirmation, persistence, and enrollment completion."""
    with app.app_context():
        officer = create_user("officer_join@test.com", "Pass@1234", "Sunil", "Verma", "reviewer")
        unauth_user = create_user("unauth_join@test.com", "Pass@1234", "Bystander", "User", "candidate")
        cand = create_user("cand_join@test.com", "Pass@1234", "Amit", "Kumar", "candidate")

        app_rec = CandidateApplication(
            application_number="CEF-TEST-JOIN-001",
            user_id=cand.id,
            created_by_id=officer.id,
            role_applied="HDD Helper",
            full_name="Amit Kumar",
            status=APPLICATION_STATUS_APPROVED,
            site_joining_status=JOINING_STATUS_NOT_JOINED,
        )
        db.session.add(app_rec)
        db.session.commit()
        app_id = str(app_rec.id)

    headers_officer = login(client, "officer_join@test.com", "Pass@1234")
    headers_unauth = login(client, "unauth_join@test.com", "Pass@1234")

    # 1. Unauthorized user cannot update joining status
    unauth_resp = client.post(
        f"/api/v1/applications/{app_id}/joining-status",
        headers=headers_unauth,
        json={"status": "JOINED", "joining_date": "2026-11-01"},
    )
    assert unauth_resp.status_code == 403

    # 2. Initial state: Not joined, enrollment NOT completed
    check_initial = client.get(f"/api/v1/applications/{app_id}", headers=headers_officer)
    assert check_initial.json["data"]["site_joining_status"] == "NOT_JOINED"
    assert check_initial.json["data"]["enrollment_completed"] is False

    # 3. Authorized user updates joining status to JOINED
    join_resp = client.post(
        f"/api/v1/applications/{app_id}/joining-status",
        headers=headers_officer,
        json={
            "status": "JOINED",
            "joining_date": "01/11/2026",
            "notes": "Verified at site office and assigned to unit 4.",
        },
    )
    assert join_resp.status_code == 200
    join_data = join_resp.json["data"]
    assert join_data["site_joining_status"] == "JOINED"
    assert join_data["site_joining_date"] == "01/11/2026"
    assert join_data["site_joining_updated_by_name"] == "Sunil Verma"
    assert join_data["site_joining_updated_at"] is not None
    # Enrollment completion is strictly linked to site joining confirmation!
    assert join_data["enrollment_completed"] is True

    # 4. Persistence verification: New login / fetch returns exact saved state
    headers_fresh = login(client, "officer_join@test.com", "Pass@1234")
    fetch_resp = client.get(f"/api/v1/applications/{app_id}", headers=headers_fresh)
    assert fetch_resp.status_code == 200
    persisted = fetch_resp.json["data"]
    assert persisted["site_joining_status"] == "JOINED"
    assert persisted["site_joining_date"] == "01/11/2026"
    assert persisted["enrollment_completed"] is True

    # 5. Correction of joining status: Authorized user can correct information
    correct_resp = client.post(
        f"/api/v1/applications/{app_id}/joining-status",
        headers=headers_officer,
        json={
            "status": "NOT_JOINED",
            "notes": "Corrected: candidate reporting deferred to next week.",
        },
    )
    assert correct_resp.status_code == 200
    assert correct_resp.json["data"]["site_joining_status"] == "NOT_JOINED"
    assert correct_resp.json["data"]["enrollment_completed"] is False


def test_excel_export_multi_worksheet_integration(client, app):
    """Test Requirement 4 & 8: Excel workbook contains Approved, Offer Letter Released, Joined On Site worksheets."""
    with app.app_context():
        admin = create_user("excel_admin@test.com", "Pass@1234", "Rajesh", "Gupta", "admin")
        cand1 = create_user("c1@test.com", "Pass@1234", "Candidate", "One", "candidate")
        cand2 = create_user("c2@test.com", "Pass@1234", "Candidate", "Two", "candidate")
        cand3 = create_user("c3@test.com", "Pass@1234", "Candidate", "Three", "candidate")

        # Candidate 1: Approved only (no offer letter, not joined)
        app1 = CandidateApplication(
            application_number="CEF-EXCEL-001",
            user_id=cand1.id,
            created_by_id=admin.id,
            role_applied="HDD Operate",
            full_name="Candidate One Approved",
            status=APPLICATION_STATUS_APPROVED,
            site_joining_status=JOINING_STATUS_NOT_JOINED,
        )
        app1.reviewed_at = datetime.now(timezone.utc)
        db.session.add(app1)

        # Candidate 2: Approved + Offer Letter Released (not joined)
        app2 = CandidateApplication(
            application_number="CEF-EXCEL-002",
            user_id=cand2.id,
            created_by_id=admin.id,
            role_applied="HDD Tracker",
            full_name="Candidate Two Offered",
            status=APPLICATION_STATUS_APPROVED,
            site_joining_status=JOINING_STATUS_NOT_JOINED,
        )
        app2.reviewed_at = datetime.now(timezone.utc)
        db.session.add(app2)
        db.session.flush()

        letter2 = CandidateLetter(
            application_id=app2.id,
            letter_type="offer_letter",
            docx_storage_path="storage/dummy2.docx",
            status=LETTER_STATUS_SENT,
            sent_at=datetime.now(timezone.utc),
            sent_by_id=admin.id,
        )
        db.session.add(letter2)

        # Candidate 3: Approved + Offer Letter Released + Joined On Site
        app3 = CandidateApplication(
            application_number="CEF-EXCEL-003",
            user_id=cand3.id,
            created_by_id=admin.id,
            role_applied="HDD Helper",
            full_name="Candidate Three Joined",
            status=APPLICATION_STATUS_APPROVED,
            site_joining_status=JOINING_STATUS_JOINED,
            site_joining_date="25/10/2026",
            site_joining_updated_by_id=admin.id,
            site_joining_updated_at=datetime.now(timezone.utc),
        )
        app3.reviewed_at = datetime.now(timezone.utc)
        db.session.add(app3)
        db.session.flush()

        letter3 = CandidateLetter(
            application_id=app3.id,
            letter_type="offer_letter",
            docx_storage_path="storage/dummy3.docx",
            status=LETTER_STATUS_SENT,
            sent_at=datetime.now(timezone.utc),
            sent_by_id=admin.id,
        )
        db.session.add(letter3)

        db.session.commit()

    headers = login(client, "excel_admin@test.com", "Pass@1234")

    # Download Excel Report
    res = client.get("/api/v1/applications/export?candidate_name=Candidate", headers=headers)
    assert res.status_code == 200
    assert "spreadsheetml.sheet" in res.content_type

    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    sheetnames = wb.sheetnames
    assert "Candidate Applications" in sheetnames
    assert "Approved" in sheetnames
    assert "Offer Letter Released" in sheetnames
    assert "Joined On Site" in sheetnames

    # Check "Approved" worksheet
    ws_appr = wb["Approved"]
    appr_rows = list(ws_appr.iter_rows(values_only=True))
    header_appr = appr_rows[0]
    assert "Candidate Application ID" in header_appr
    assert "Candidate Name" in header_appr
    assert "Approval Status" in header_appr
    assert "Offer Letter Status" in header_appr
    assert "Site Joining Status" in header_appr
    # Must contain candidates 1, 2, and 3 (all 3 are approved)
    appr_names = [r[1] for r in appr_rows[1:]]
    assert "Candidate One Approved" in appr_names
    assert "Candidate Two Offered" in appr_names
    assert "Candidate Three Joined" in appr_names

    # Check "Offer Letter Released" worksheet
    ws_offer = wb["Offer Letter Released"]
    offer_rows = list(ws_offer.iter_rows(values_only=True))
    offer_names = [r[1] for r in offer_rows[1:]]
    # Candidate 1 must NOT appear in Offer Letter Released
    assert "Candidate One Approved" not in offer_names
    # Candidates 2 and 3 must appear
    assert "Candidate Two Offered" in offer_names
    assert "Candidate Three Joined" in offer_names

    # Check "Joined On Site" worksheet
    ws_joined = wb["Joined On Site"]
    joined_rows = list(ws_joined.iter_rows(values_only=True))
    joined_names = [r[1] for r in joined_rows[1:]]
    # Only Candidate 3 has site_joining_status == JOINED
    assert "Candidate One Approved" not in joined_names
    assert "Candidate Two Offered" not in joined_names
    assert "Candidate Three Joined" in joined_names

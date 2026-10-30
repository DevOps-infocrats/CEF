import io
import os
import hashlib
import uuid
import docx
import pytest
from app.extensions import db
from app.modules.auth.models import User, Role
from app.modules.auth.services import hash_password
from app.modules.candidates.models import (
    CandidateApplication,
    CandidateDocument,
    APPLICATION_STATUS_APPROVED,
)
from app.modules.letters.models import (
    CandidateLetter,
    DocumentTemplate,
    TEMPLATE_TYPE_LOI,
    TEMPLATE_TYPE_OFFER,
)
from app.modules.letters.services import LetterTemplateService, resolve_storage_path


def compute_sha256(content: bytes) -> str:
    """Compute SHA-256 hex digest of raw byte content."""
    return hashlib.sha256(content).hexdigest()


def create_sample_docx(marker_text: str) -> bytes:
    """Generate a valid DOCX file in memory containing specific markers and table data."""
    doc = docx.Document()
    doc.add_heading(f"Clean Energy Fund Master Template — {marker_text}", level=0)
    doc.add_paragraph(f"This is an official template with unique marker: {marker_text}")
    doc.add_paragraph("Candidate Name: {{candidate_name}}")
    doc.add_paragraph("Designation: {{designation}}")
    doc.add_paragraph("Joining Date: {{joining_date}}")
    doc.add_paragraph("Monthly CTC: {{monthly_ctc}}")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Candidate Reference"
    table.cell(1, 1).text = "{{application_number}}"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_user_with_role(email: str, password: str, role_name: str) -> User:
    user = User(
        email=email,
        username=email.split("@")[0],
        password_hash=hash_password(password),
        first_name="Admin",
        last_name="Tester",
        is_active=True,
        is_verified=True,
        approval_status="APPROVED",
    )
    role = Role.query.filter_by(name=role_name).first()
    if role:
        user.roles.append(role)
    db.session.add(user)
    db.session.commit()
    return user


def login_client(client, email: str, password: str) -> dict:
    resp = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert resp.status_code == 200, f"Login failed for {email}: {resp.json}"
    return {"Authorization": f"Bearer {resp.json['data']['access_token']}"}


def test_master_template_exact_byte_preservation_and_replacement(client, app):
    """
    Verify that uploaded master documents:
    1. Are saved directly without modification.
    2. Download byte-for-byte identical with matching SHA-256 cryptographic hashes.
    3. Replacement activates new version and downloads the new file byte-for-byte.
    4. Candidate letter generation never alters the stored master template.
    """
    with app.app_context():
        admin = create_user_with_role("master_admin@example.com", "AdminPass@123", "admin")

    headers = login_client(client, "master_admin@example.com", "AdminPass@123")

    # 1. Create original template V1 bytes
    original_v1_bytes = create_sample_docx("ORIGINAL_TEMPLATE_V1_TEST_MARKER_BYTE_FOR_BYTE")
    original_v1_hash = compute_sha256(original_v1_bytes)

    # 2. Upload master template
    upload_data = {
        "file": (io.BytesIO(original_v1_bytes), "Custom_Master_LOI_v1.docx"),
    }
    upload_res = client.post(
        "/api/v1/letters/templates/letter_of_intent/upload",
        data=upload_data,
        content_type="multipart/form-data",
        headers=headers,
    )
    assert upload_res.status_code == 201
    upload_json = upload_res.get_json()
    assert upload_json["success"] is True
    template_id = upload_json["data"]["id"]
    version_1 = upload_json["data"]["version"]

    # 3. Verify disk file bytes directly match original upload bytes
    with app.app_context():
        tmpl_rec = db.session.get(DocumentTemplate, uuid.UUID(template_id))
        assert tmpl_rec is not None
        assert tmpl_rec.is_active is True
        resolved_path = resolve_storage_path(tmpl_rec.storage_path)
        assert os.path.exists(resolved_path)

        with open(resolved_path, "rb") as f:
            disk_bytes = f.read()
        assert compute_sha256(disk_bytes) == original_v1_hash, "Stored bytes on disk do not match original upload hash!"

    # 4. Download master template via API and verify downloaded bytes match original upload exactly
    download_res = client.get(
        "/api/v1/letters/templates/letter_of_intent/download",
        headers=headers,
    )
    assert download_res.status_code == 200
    assert download_res.headers.get("Content-Type") == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    downloaded_bytes = download_res.data
    downloaded_hash = compute_sha256(downloaded_bytes)

    assert downloaded_hash == original_v1_hash, (
        f"Byte mismatch! Original SHA-256: {original_v1_hash}, Downloaded SHA-256: {downloaded_hash}"
    )
    assert len(downloaded_bytes) == len(original_v1_bytes)

    # 5. Master Template Replacement: Upload a completely distinct second master template
    original_v2_bytes = create_sample_docx("REPLACEMENT_TEMPLATE_V2_WITH_NEW_STYLES_AND_TABLES")
    original_v2_hash = compute_sha256(original_v2_bytes)
    assert original_v2_hash != original_v1_hash

    upload_v2_data = {
        "file": (io.BytesIO(original_v2_bytes), "Custom_Master_LOI_v2.docx"),
    }
    upload_v2_res = client.post(
        "/api/v1/letters/templates/letter_of_intent/upload",
        data=upload_v2_data,
        content_type="multipart/form-data",
        headers=headers,
    )
    assert upload_v2_res.status_code == 201
    upload_v2_json = upload_v2_res.get_json()
    version_2 = upload_v2_json["data"]["version"]
    assert version_2 > version_1

    # 6. Verify subsequent download immediately serves the newly uploaded replacement byte-for-byte
    download_v2_res = client.get(
        "/api/v1/letters/templates/letter_of_intent/download",
        headers=headers,
    )
    assert download_v2_res.status_code == 200
    downloaded_v2_bytes = download_v2_res.data
    downloaded_v2_hash = compute_sha256(downloaded_v2_bytes)

    assert downloaded_v2_hash == original_v2_hash, (
        f"Replacement download failed! Expected SHA-256: {original_v2_hash}, Got: {downloaded_v2_hash}"
    )

    # 7. Candidate generation must NEVER mutate or touch the active master template
    with app.app_context():
        active_tmpl = LetterTemplateService.get_active_template(TEMPLATE_TYPE_LOI)
        master_path = resolve_storage_path(active_tmpl.storage_path)
        with open(master_path, "rb") as f:
            master_bytes_before = f.read()
        master_hash_before = compute_sha256(master_bytes_before)

        # Create candidate application
        cand = create_user_with_role("cand_letter_test@example.com", "CandPass@123", "candidate")
        app_record = CandidateApplication(
            user_id=cand.id,
            application_number="CEF-TEST-BYTE-001",
            full_name="Byte Verification Candidate",
            email="cand_letter_test@example.com",
            mobile_number="9876543210",
            role_applied="Field Engineer",
            negotiated_salary=45000,
            status=APPLICATION_STATUS_APPROVED,
        )
        db.session.add(app_record)
        db.session.commit()
        app_id_str = str(app_record.id)

    # Generate candidate letters
    letters_res = client.get(f"/api/v1/applications/{app_id_str}/letters", headers=headers)
    assert letters_res.status_code == 200

    # Verify candidate letters were generated separately and master template remains 100% byte-for-byte identical
    with open(master_path, "rb") as f:
        master_bytes_after = f.read()
    master_hash_after = compute_sha256(master_bytes_after)

    assert master_hash_after == master_hash_before, "Master template was modified during candidate letter generation!"
    assert master_hash_after == original_v2_hash, "Master template SHA-256 diverged from original uploaded bytes!"


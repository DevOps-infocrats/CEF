import io
import uuid
from decimal import Decimal
from datetime import datetime, timezone
from app.extensions import db
from app.modules.auth.models import (
    User,
    Role,
    UserRole,
    UserPermission,
    Permission,
)
from app.modules.auth.services import hash_password
from app.modules.candidates.models import (
    CandidateApplication,
    APPLICATION_STATUS_APPROVED,
    APPLICATION_STATUS_PENDING,
)
from app.modules.letters.models import (
    DocumentTemplate,
    CandidateLetter,
    TEMPLATE_TYPE_LOI,
    TEMPLATE_TYPE_OFFER,
    LETTER_STATUS_DRAFT,
    LETTER_STATUS_SENT,
)


def create_user(email: str, password: str, permissions: list[str] = None, role: str = "staff"):
    user = User(
        email=email.lower().strip(),
        username=email.split("@")[0],
        password_hash=hash_password(password),
        first_name="Test",
        last_name="User",
        is_active=True,
        is_verified=True,
        approval_status="APPROVED",
    )
    db.session.add(user)
    db.session.flush()

    r = Role.query.filter_by(name=role).first()
    if r:
        ur = UserRole(user_id=user.id, role_id=r.id)
        db.session.add(ur)

    if permissions:
        for p_code in permissions:
            perm = Permission.query.filter_by(code=p_code).first()
            if perm:
                up = UserPermission(user_id=user.id, permission_id=perm.id)
                db.session.add(up)

    db.session.commit()
    return user


def login(client, email: str, password: str):
    resp = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert resp.status_code == 200, f"Login failed for {email}: {resp.json}"
    return {"Authorization": f"Bearer {resp.json['data']['access_token']}"}


def create_test_application(user: User, status: str = APPLICATION_STATUS_APPROVED) -> CandidateApplication:
    app = CandidateApplication(
        application_number=f"APP-{uuid.uuid4().hex[:8].upper()}",
        user_id=user.id,
        hero_category="HDD",
        role_applied="HDD Operator",
        full_name="Rajesh Kumar Sharma",
        mobile_number="9876543210",
        email="rajesh.kumar@example.com",
        aadhaar_number="987654321012",
        pan_number="ABCDE1234F",
        current_address_line1="123 Civil Lines",
        current_city="Jaipur",
        current_state="Rajasthan",
        current_pincode="302001",
        permanent_address_line1="123 Civil Lines",
        permanent_city="Jaipur",
        permanent_state="Rajasthan",
        permanent_pincode="302001",
        account_holder_name="Rajesh Kumar",
        bank_name="State Bank of India",
        account_number="12345678901",
        ifsc_code="SBIN0001234",
        circle="Rajasthan",
        district="Jaipur",
        block="Amber",
        current_salary=Decimal("35000.00"),
        asking_salary=Decimal("42000.00"),
        negotiated_salary=Decimal("40000.00"),
        status=status,
    )
    db.session.add(app)
    db.session.commit()
    return app


def test_master_templates_metadata_and_download(client):
    user = create_user("tmpl_viewer@example.com", "Pass@1234", ["portal:access", "users:manage"])
    headers = login(client, "tmpl_viewer@example.com", "Pass@1234")

    # 1. Fetch Letter of Intent metadata
    resp_loi = client.get("/api/v1/letters/templates/letter_of_intent", headers=headers)
    assert resp_loi.status_code == 200
    loi_data = resp_loi.json["data"]
    assert loi_data["template_type"] == "letter_of_intent"
    assert loi_data["version"] >= 1
    assert loi_data["is_active"] is True
    assert len(loi_data["field_mappings"]) > 0

    # 2. Fetch Offer Letter metadata
    resp_offer = client.get("/api/v1/letters/templates/offer_letter", headers=headers)
    assert resp_offer.status_code == 200
    offer_data = resp_offer.json["data"]
    assert offer_data["template_type"] == "offer_letter"
    assert offer_data["version"] >= 1
    assert offer_data["is_active"] is True

    # 3. Download master template
    dl_resp = client.get("/api/v1/letters/templates/letter_of_intent/download", headers=headers)
    assert dl_resp.status_code == 200
    assert "wordprocessingml" in dl_resp.content_type


def test_admin_upload_new_master_template(client):
    import docx
    admin = create_user("admin_tmpl@example.com", "Pass@1234", ["users:manage"], role="admin")
    regular = create_user("regular_user@example.com", "Pass@1234", ["portal:access"], role="staff")

    admin_headers = login(client, "admin_tmpl@example.com", "Pass@1234")
    regular_headers = login(client, "regular_user@example.com", "Pass@1234")

    # 1. Corrupted file upload returns 400
    resp_corrupt = client.post(
        "/api/v1/letters/templates/letter_of_intent/upload",
        headers=admin_headers,
        data={"file": (io.BytesIO(b"fake data"), "broken.docx")},
        content_type="multipart/form-data",
    )
    assert resp_corrupt.status_code == 400

    # 2. Regular user cannot upload
    resp_fail = client.post(
        "/api/v1/letters/templates/letter_of_intent/upload",
        headers=regular_headers,
        data={"file": (io.BytesIO(b"dummy"), "new_loi.docx")},
        content_type="multipart/form-data",
    )
    assert resp_fail.status_code == 403

    # 3. Admin uploads valid new template
    valid_doc = docx.Document()
    valid_doc.add_heading("Letter of Intent Master v2", 0)
    valid_doc.add_paragraph("Dear {{candidate_full_name}},\nJoining Date: {{joining_date}}")
    buf = io.BytesIO()
    valid_doc.save(buf)
    buf.seek(0)

    resp_ok = client.post(
        "/api/v1/letters/templates/letter_of_intent/upload",
        headers=admin_headers,
        data={"file": (buf, "revised_loi_master.docx")},
        content_type="multipart/form-data",
    )
    assert resp_ok.status_code == 201
    assert resp_ok.json["data"]["version"] >= 2
    assert resp_ok.json["data"]["is_active"] is True


def test_candidate_letters_generation_and_draft_editing(client):
    approver = create_user("approver_doc@example.com", "Pass@1234", ["candidates:approve", "documents:view"])
    headers = login(client, "approver_doc@example.com", "Pass@1234")

    # 1. Blocked for unapproved candidate
    cand_user = create_user("cand_pending@example.com", "Pass@1234", ["portal:access"], role="candidate")
    pending_app = create_test_application(cand_user, status=APPLICATION_STATUS_PENDING)

    resp_pending = client.get(f"/api/v1/applications/{pending_app.id}/letters", headers=headers)
    assert resp_pending.status_code == 400
    assert "APPROVED" in resp_pending.json["message"]

    # 2. Allowed for approved candidate
    approved_app = create_test_application(cand_user, status=APPLICATION_STATUS_APPROVED)
    resp_gen = client.get(f"/api/v1/applications/{approved_app.id}/letters", headers=headers)
    assert resp_gen.status_code == 200
    letters = resp_gen.json["data"]
    assert letters["application"]["full_name"] == "Rajesh Kumar Sharma"
    assert "letter_of_intent" in letters
    assert "offer_letter" in letters
    assert letters["letter_of_intent"]["status"] == LETTER_STATUS_DRAFT

    # 3. Update joining date and monthly salary in draft
    patch_resp = client.patch(
        f"/api/v1/applications/{approved_app.id}/letters/letter_of_intent",
        headers=headers,
        json={
            "field_values": {
                "joining_date": "15/11/2026",
                "monthly_salary": "42000",
            }
        },
    )
    assert patch_resp.status_code == 200
    updated_fields = patch_resp.json["data"]["field_values"]
    assert updated_fields["joining_date"] == "15/11/2026"
    assert updated_fields["monthly_salary"] == "42000"

    # Also update offer letter joining date
    patch_offer = client.patch(
        f"/api/v1/applications/{approved_app.id}/letters/offer_letter",
        headers=headers,
        json={
            "field_values": {
                "joining_date": "15/11/2026",
                "monthly_salary": "42000",
            }
        },
    )
    assert patch_offer.status_code == 200

    # 4. Download candidate docx
    dl_resp = client.get(f"/api/v1/applications/{approved_app.id}/letters/letter_of_intent/download?format=docx", headers=headers)
    assert dl_resp.status_code == 200
    assert len(dl_resp.data) > 0


def test_send_letters_workflow_and_duplicate_prevention(client):
    approver = create_user("approver_sender@example.com", "Pass@1234", ["candidates:approve", "documents:view"])
    headers = login(client, "approver_sender@example.com", "Pass@1234")

    cand_user = create_user("cand_send@example.com", "Pass@1234", ["portal:access"], role="candidate")
    approved_app = create_test_application(cand_user, status=APPLICATION_STATUS_APPROVED)

    # Initialize letters
    client.get(f"/api/v1/applications/{approved_app.id}/letters", headers=headers)

    # 1. Attempt send without joining date -> blocks
    send_fail = client.post(f"/api/v1/applications/{approved_app.id}/letters/send", headers=headers, json={})
    assert send_fail.status_code == 400
    assert "Joining Date" in send_fail.json["message"]

    # 2. Set joining date for both
    client.patch(
        f"/api/v1/applications/{approved_app.id}/letters/letter_of_intent",
        headers=headers,
        json={"field_values": {"joining_date": "10/11/2026"}},
    )
    client.patch(
        f"/api/v1/applications/{approved_app.id}/letters/offer_letter",
        headers=headers,
        json={"field_values": {"joining_date": "10/11/2026"}},
    )

    # 3. Successful send
    send_ok = client.post(f"/api/v1/applications/{approved_app.id}/letters/send", headers=headers, json={})
    assert send_ok.status_code == 200
    assert send_ok.json["data"]["success"] is True
    assert send_ok.json["data"]["recipient"] == "rajesh.kumar@example.com"

    # Verify letter records in DB are marked SENT
    letters_in_db = CandidateLetter.query.filter_by(application_id=approved_app.id).all()
    assert len(letters_in_db) == 2
    for l in letters_in_db:
        assert l.status == LETTER_STATUS_SENT
        assert l.sent_at is not None
        assert l.email_status == "SENT"

    # 4. Attempt duplicate send -> blocked without force
    dup_resp = client.post(f"/api/v1/applications/{approved_app.id}/letters/send", headers=headers, json={})
    assert dup_resp.status_code == 400
    assert "already been sent" in dup_resp.json["message"]

    # 5. Non-approver user cannot send
    unauthorized = create_user("no_send_perm@example.com", "Pass@1234", ["portal:access", "candidates:read"])
    unauth_headers = login(client, "no_send_perm@example.com", "Pass@1234")
    unauth_send = client.post(f"/api/v1/applications/{approved_app.id}/letters/send", headers=unauth_headers, json={})
    assert unauth_send.status_code == 403


def test_click_to_edit_offer_letter_text_persistence(client):
    """Test click-to-edit Offer Letter changes across page layout persist to candidate copy and reopen intact."""
    user = create_user("editor_officer@example.com", "Pass@1234", ["candidates:approve", "documents:view"])
    headers = login(client, "editor_officer@example.com", "Pass@1234")

    cand_user = create_user("cand_edit@example.com", "Pass@1234", ["portal:access"], role="candidate")
    approved_app = create_test_application(cand_user, status=APPLICATION_STATUS_APPROVED)

    # 1. Fetch letters to initialize candidate draft
    resp = client.get(f"/api/v1/applications/{approved_app.id}/letters", headers=headers)
    assert resp.status_code == 200
    orig_offer = resp.json["data"]["offer_letter"]

    # 2. Save click-to-edit text edits across the document
    custom_edits = {
        "reference_number": "FFPL/CUSTOM/2026/99999",
        "current_date": "26th Oct, 2026",
        "candidate_full_name": "Rajesh Custom Sharma",
        "candidate_mobile": "9888877777",
        "approved_role_name": "Lead Optical Engineer",
        "approved_location": "Jaipur Tech Hub",
        "joining_date": "05/11/2026",
        "custom_content": {
            "subject": "Re: Official Offer Letter - Lead Optical Engineer",
            "salutation": "Dear Rajesh Custom Sharma,",
            "p10": "You will be happy to know that you have been appointed as a Lead Optical Engineer in our Organization. Presently you will be posted in Jaipur Tech Hub.",
            "p11": "Special performance reviews will be conducted quarterly.",
            "yours_faithfully": "Warm regards and best wishes,",
            "hr_dept": "Central HR & Talent Acquisition",
            "company_name": "For, Fiber Force Private Limited (Corporate Division)",
            "annex1_s1_h": "Service subject to customized company rules",
            "annex2_title": "Annexure – II: Verified CTC Breakdown",
        },
        "custom_paragraphs": {
            "subject": "Re: Official Offer Letter - Lead Optical Engineer",
            "salutation": "Dear Rajesh Custom Sharma,",
            "p10": "You will be happy to know that you have been appointed as a Lead Optical Engineer in our Organization. Presently you will be posted in Jaipur Tech Hub.",
            "p11": "Special performance reviews will be conducted quarterly.",
            "yours_faithfully": "Warm regards and best wishes,",
            "hr_dept": "Central HR & Talent Acquisition",
            "company_name": "For, Fiber Force Private Limited (Corporate Division)",
            "annex1_s1_h": "Service subject to customized company rules",
            "annex2_title": "Annexure – II: Verified CTC Breakdown",
        },
    }

    patch_resp = client.patch(
        f"/api/v1/applications/{approved_app.id}/letters/offer_letter",
        headers=headers,
        json={"field_values": custom_edits},
    )
    assert patch_resp.status_code == 200
    patched_data = patch_resp.json["data"]
    assert patched_data["field_values"]["reference_number"] == "FFPL/CUSTOM/2026/99999"
    assert patched_data["field_values"]["custom_content"]["subject"] == "Re: Official Offer Letter - Lead Optical Engineer"

    # 3. Simulate reopening the Offer Letter preview (re-fetching candidate letters)
    reopen_resp = client.get(f"/api/v1/applications/{approved_app.id}/letters", headers=headers)
    assert reopen_resp.status_code == 200
    reopened_offer = reopen_resp.json["data"]["offer_letter"]
    reopened_fields = reopened_offer["field_values"]

    # Verify changes remain intact upon reopening
    assert reopened_fields["reference_number"] == "FFPL/CUSTOM/2026/99999"
    assert reopened_fields["current_date"] == "26th Oct, 2026"
    assert reopened_fields["candidate_full_name"] == "Rajesh Custom Sharma"
    assert reopened_fields["custom_content"]["subject"] == "Re: Official Offer Letter - Lead Optical Engineer"
    assert reopened_fields["custom_content"]["yours_faithfully"] == "Warm regards and best wishes,"
    assert reopened_fields["custom_content"]["p10"].startswith("You will be happy to know that you have been appointed as a Lead Optical Engineer")

    # 4. Verify candidate-specific docx is generated and downloadable
    cand_docx_resp = client.get(
        f"/api/v1/applications/{approved_app.id}/letters/offer_letter/download?format=docx",
        headers=headers,
    )
    assert cand_docx_resp.status_code == 200
    assert len(cand_docx_resp.data) > 0

    # 5. Verify master template is untouched
    master_resp = client.get("/api/v1/letters/templates/offer_letter", headers=headers)
    assert master_resp.status_code == 200
    assert master_resp.json["data"]["template_type"] == "offer_letter"


def test_unapproved_candidate_cannot_generate_letters(client):
    """Ensure candidate letters can strictly only be generated for APPROVED candidates."""
    user = create_user("verifier_guard@test.com", "Secret123!", permissions=["letters:view", "letters:generate", "candidate:view"])
    headers = login(client, "verifier_guard@test.com", "Secret123!")

    # Candidate in PENDING status
    pending_app = create_test_application(user, status=APPLICATION_STATUS_PENDING)

    resp = client.get(f"/api/v1/applications/{pending_app.id}/letters", headers=headers)
    assert resp.status_code == 400
    assert "Letters can only be generated for APPROVED candidates" in resp.json["message"]


def test_nonexistent_candidate_letters_returns_404(client):
    """Ensure invalid candidate ID returns 404 error."""
    user = create_user("verifier_404@test.com", "Secret123!", permissions=["letters:view", "letters:generate", "candidate:view"])
    headers = login(client, "verifier_404@test.com", "Secret123!")

    random_id = uuid.uuid4()
    resp = client.get(f"/api/v1/applications/{random_id}/letters", headers=headers)
    assert resp.status_code == 404


def test_approved_candidate_letters_duplicate_resilience(client):
    """Ensure multiple calls or duplicate rows do not cause MultipleResultsFound or failure."""
    user = create_user("verifier_dup@test.com", "Secret123!", permissions=["letters:view", "letters:generate", "candidate:view"])
    headers = login(client, "verifier_dup@test.com", "Secret123!")

    approved_app = create_test_application(user, status=APPLICATION_STATUS_APPROVED)

    # First call generates letters
    resp1 = client.get(f"/api/v1/applications/{approved_app.id}/letters", headers=headers)
    assert resp1.status_code == 200
    assert resp1.json["data"]["application"]["status"] == "APPROVED"

    # Second concurrent or repeated call should also succeed cleanly
    resp2 = client.get(f"/api/v1/applications/{approved_app.id}/letters", headers=headers)
    assert resp2.status_code == 200

    # Inline download calls should succeed cleanly without MultipleResultsFound
    dl_resp = client.get(f"/api/v1/applications/{approved_app.id}/letters/offer_letter/download?format=pdf&inline=1", headers=headers)
    assert dl_resp.status_code == 200
    assert dl_resp.mimetype == "application/pdf"


def test_one_master_template_reused_across_multiple_candidates(client):
    """Verify one master template is reused for multiple candidates, populating real data without modifying the master."""
    import hashlib
    import os

    admin = create_user("admin_multi@example.com", "Pass@1234", ["users:manage", "letters:view", "letters:generate", "candidate:view"], role="admin")
    admin_headers = login(client, "admin_multi@example.com", "Pass@1234")

    # 1. Fetch currently active master template and compute its hash
    from app.modules.letters.services import LetterTemplateService
    active_master = LetterTemplateService.get_active_template("letter_of_intent")
    master_path = active_master.storage_path
    assert os.path.exists(master_path)
    with open(master_path, "rb") as f:
        master_hash_before = hashlib.sha256(f.read()).hexdigest()

    # 2. Create Candidate 1 (Alice)
    u1 = create_user("alice@candidate.com", "Pass@1234", role="candidate")
    app1 = CandidateApplication(
        application_number="APP-ALICE-001",
        user_id=u1.id,
        hero_category="Hero 1",
        role_applied="Fiber Splicer",
        full_name="Alice Candidate Sharma",
        mobile_number="9111111111",
        email="alice@candidate.com",
        aadhaar_number="111122223333",
        pan_number="ABCDE1111A",
        current_address_line1="10 Alwar Road",
        current_city="Alwar",
        current_state="Rajasthan",
        current_pincode="301001",
        permanent_address_line1="10 Alwar Road",
        permanent_city="Alwar",
        permanent_state="Rajasthan",
        permanent_pincode="301001",
        account_holder_name="Alice Candidate",
        bank_name="HDFC Bank",
        account_number="12345678901",
        ifsc_code="HDFC0001234",
        circle="Rajasthan",
        district="Alwar",
        block="Alwar Central",
        current_salary=Decimal("40000"),
        asking_salary=Decimal("45000"),
        negotiated_salary=Decimal("45000"),
        status=APPLICATION_STATUS_APPROVED,
    )
    db.session.add(app1)

    # 3. Create Candidate 2 (Bob)
    u2 = create_user("bob@candidate.com", "Pass@1234", role="candidate")
    app2 = CandidateApplication(
        application_number="APP-BOB-002",
        user_id=u2.id,
        hero_category="Hero 2",
        role_applied="HDD Tracker",
        full_name="Bob Candidate Verma",
        mobile_number="9222222222",
        email="bob@candidate.com",
        aadhaar_number="444455556666",
        pan_number="ABCDE2222B",
        current_address_line1="25 Kota Road",
        current_city="Kota",
        current_state="Rajasthan",
        current_pincode="324001",
        permanent_address_line1="25 Kota Road",
        permanent_city="Kota",
        permanent_state="Rajasthan",
        permanent_pincode="324001",
        account_holder_name="Bob Candidate",
        bank_name="ICICI Bank",
        account_number="98765432101",
        ifsc_code="ICIC0005678",
        circle="Rajasthan",
        district="Kota",
        block="Kota South",
        current_salary=Decimal("30000"),
        asking_salary=Decimal("35000"),
        negotiated_salary=Decimal("35000"),
        status=APPLICATION_STATUS_APPROVED,
    )
    db.session.add(app2)
    db.session.commit()

    # 4. Generate letters for Candidate 1 from master template
    resp1 = client.get(f"/api/v1/applications/{app1.id}/letters", headers=admin_headers)
    assert resp1.status_code == 200
    loi1 = resp1.json["data"]["letter_of_intent"]
    assert loi1["field_values"]["candidate_full_name"] == "Alice Candidate Sharma"
    assert loi1["field_values"]["approved_role_name"] == "Fiber Splicer"
    assert loi1["field_values"]["monthly_salary"] == "45000"

    # 5. Generate letters for Candidate 2 from same master template
    resp2 = client.get(f"/api/v1/applications/{app2.id}/letters", headers=admin_headers)
    assert resp2.status_code == 200
    loi2 = resp2.json["data"]["letter_of_intent"]
    assert loi2["field_values"]["candidate_full_name"] == "Bob Candidate Verma"
    assert loi2["field_values"]["approved_role_name"] == "HDD Tracker"
    assert loi2["field_values"]["monthly_salary"] == "35000"

    # 6. Re-check Candidate 1 letters: ensure Candidate 2's generation did not affect Candidate 1
    resp1_check = client.get(f"/api/v1/applications/{app1.id}/letters", headers=admin_headers)
    assert resp1_check.status_code == 200
    assert resp1_check.json["data"]["letter_of_intent"]["field_values"]["candidate_full_name"] == "Alice Candidate Sharma"

    # 7. Verify the Master Template file content has NOT changed (byte-for-byte read-only)
    with open(master_path, "rb") as f:
        master_hash_after = hashlib.sha256(f.read()).hexdigest()
    assert master_hash_before == master_hash_after, "Master template was modified during candidate generation!"


def test_in_place_editing_persists_to_candidate_docx_and_pdf(client):
    """Verify that editing text directly in place saves to candidate DOCX/PDF and does not leak or alter master."""
    import hashlib
    import os
    import docx
    from app.modules.letters.services import LetterTemplateService

    admin = create_user("editor_admin@example.com", "Pass@1234", ["users:manage", "letters:view", "letters:generate", "candidates:approve", "candidates:edit"], role="admin")
    admin_headers = login(client, "editor_admin@example.com", "Pass@1234")

    # Record master SHA256 before
    master_tmpl = LetterTemplateService.get_active_template("letter_of_intent")
    with open(master_tmpl.storage_path, "rb") as f:
        master_hash_before = hashlib.sha256(f.read()).hexdigest()

    # Create Candidate A
    uA = create_user("cand_a@test.com", "Pass@1234", role="candidate")
    appA = create_test_application(uA, status=APPLICATION_STATUS_APPROVED)

    # 1. Generate letters for Candidate A
    respA = client.get(f"/api/v1/applications/{appA.id}/letters", headers=admin_headers)
    assert respA.status_code == 200
    loi_data = respA.json["data"]["letter_of_intent"]
    assert "content" in loi_data
    assert "paragraphs" in loi_data["content"]

    # 2. Perform in-place text edits on Candidate A's LOI:
    # E.g. change appointment role paragraph or custom text
    custom_text = "SPECIAL CUSTOM CLAUSE: Candidate must report to headquarters on day one."
    edit_payload = {
        "field_values": {
            "joining_date": "15/11/2026",
            "monthly_salary": "48000",
        },
        "text_edits": {
            "HDD Operator": "Lead HDD Telecom Specialist",
        },
        "paragraphs": {
            "10": f"We are pleased to inform you that Fiber Force Pvt Ltd intends to engage your services as Lead HDD Telecom Specialist. {custom_text}",
        },
    }
    patch_resp = client.patch(
        f"/api/v1/applications/{appA.id}/letters/letter_of_intent",
        headers=admin_headers,
        json=edit_payload,
    )
    assert patch_resp.status_code == 200
    updated_loi = patch_resp.json["data"]
    assert updated_loi["field_values"]["joining_date"] == "15/11/2026"
    assert updated_loi["field_values"]["monthly_salary"] == "48000"

    # 3. Verify Candidate A's DOCX on disk contains the exact custom edited text
    cand_a_letter = CandidateLetter.query.filter_by(application_id=appA.id, letter_type="letter_of_intent").order_by(CandidateLetter.version.desc()).first()
    assert os.path.exists(cand_a_letter.docx_storage_path)
    cand_a_doc = docx.Document(cand_a_letter.docx_storage_path)
    full_doc_text = "\n".join(p.text for p in cand_a_doc.paragraphs)
    assert custom_text in full_doc_text
    assert "Lead HDD Telecom Specialist" in full_doc_text

    # 4. Create Candidate B and generate their letter
    uB = create_user("cand_b@test.com", "Pass@1234", role="candidate")
    appB = create_test_application(uB, status=APPLICATION_STATUS_APPROVED)
    appB.full_name = "Candidate B Unique"
    appB.role_applied = "Fiber Technician"
    db.session.commit()

    respB = client.get(f"/api/v1/applications/{appB.id}/letters", headers=admin_headers)
    assert respB.status_code == 200
    loi_b = respB.json["data"]["letter_of_intent"]

    # 5. Verify Candidate B does NOT contain Candidate A's custom text
    cand_b_letter = CandidateLetter.query.filter_by(application_id=appB.id, letter_type="letter_of_intent").order_by(CandidateLetter.version.desc()).first()
    cand_b_doc = docx.Document(cand_b_letter.docx_storage_path)
    b_doc_text = "\n".join(p.text for p in cand_b_doc.paragraphs)
    assert custom_text not in b_doc_text, "Candidate A's custom text leaked into Candidate B's letter!"

    # 6. Re-fetch Candidate A: ensure custom edits persisted
    refetch_a = client.get(f"/api/v1/applications/{appA.id}/letters", headers=admin_headers)
    assert refetch_a.status_code == 200
    a_content = refetch_a.json["data"]["letter_of_intent"]["content"]
    assert custom_text in a_content["paragraphs"].get("10", "")

    # 7. Verify the Master Template file content remains 100% untouched
    with open(master_tmpl.storage_path, "rb") as f:
        master_hash_after = hashlib.sha256(f.read()).hexdigest()
    assert master_hash_before == master_hash_after, "Master template was modified!"


def test_offer_letter_in_place_editing_persists_and_isolates(client):
    """Verify that in-place editing of Offer Letter appointment details & Annexure II table saves to candidate DOCX and does not modify master."""
    import hashlib
    import os
    import docx
    from app.modules.letters.services import LetterTemplateService

    admin = create_user("offer_editor@example.com", "Pass@1234", ["users:manage", "letters:view", "letters:generate", "candidates:approve", "candidates:edit"], role="admin")
    admin_headers = login(client, "offer_editor@example.com", "Pass@1234")

    master_offer = LetterTemplateService.get_active_template("offer_letter")
    with open(master_offer.storage_path, "rb") as f:
        master_hash_before = hashlib.sha256(f.read()).hexdigest()

    u = create_user("cand_offer@test.com", "Pass@1234", role="candidate")
    app = create_test_application(u, status=APPLICATION_STATUS_APPROVED)

    # 1. Generate Offer Letter
    resp = client.get(f"/api/v1/applications/{app.id}/letters", headers=admin_headers)
    assert resp.status_code == 200
    offer_data = resp.json["data"]["offer_letter"]
    assert "content" in offer_data
    assert "tables" in offer_data["content"]
    assert len(offer_data["content"]["tables"]) > 0

    # 2. Patch with table cell edits and joining date
    patch_payload = {
        "field_values": {
            "joining_date": "05/12/2026",
            "monthly_salary": "52000",
            "reference_number": "FFPL/CUSTOM/REF/99999",
        },
        "table_cells": {
            "0,3,1": "52000",
            "0,8,1": "31200",
        },
        "text_edits": {
            "HDD Operator": "Senior Project Manager",
        },
    }
    patch_resp = client.patch(
        f"/api/v1/applications/{app.id}/letters/offer_letter",
        headers=admin_headers,
        json=patch_payload,
    )
    assert patch_resp.status_code == 200
    updated = patch_resp.json["data"]
    assert updated["field_values"]["joining_date"] == "05/12/2026"
    assert updated["field_values"]["monthly_salary"] == "52000"

    # 3. Verify DOCX on disk
    letter_obj = CandidateLetter.query.filter_by(application_id=app.id, letter_type="offer_letter").order_by(CandidateLetter.version.desc()).first()
    assert os.path.exists(letter_obj.docx_storage_path)
    cand_doc = docx.Document(letter_obj.docx_storage_path)
    full_text = "\n".join(p.text for p in cand_doc.paragraphs)
    assert "Senior Project Manager" in full_text
    assert "FFPL/CUSTOM/REF/99999" in full_text

    # Verify table in DOCX has updated cells
    tbl = cand_doc.tables[0]
    assert tbl.rows[3].cells[1].text == "52000"
    assert tbl.rows[8].cells[1].text == "31200"

    # 4. Verify master template is untouched
    with open(master_offer.storage_path, "rb") as f:
        master_hash_after = hashlib.sha256(f.read()).hexdigest()
    assert master_hash_before == master_hash_after, "Master Offer Letter template was altered!"



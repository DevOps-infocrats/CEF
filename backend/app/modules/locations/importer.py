import io
import openpyxl
from typing import Dict, Any, Union
from sqlalchemy import select, func
from app.extensions import db
from app.modules.locations.models import LocationState, LocationDistrict, LocationBlock


def parse_location_workbook(file_bytes_or_path: Union[bytes, str], default_state: str = "Rajasthan") -> Dict[str, Any]:
    """Parse Excel workbook containing District Summary and/or Block Mapping."""
    if isinstance(file_bytes_or_path, (bytes, bytearray)):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes_or_path), data_only=True)
    else:
        wb = openpyxl.load_workbook(file_bytes_or_path, data_only=True)

    summary_counts: Dict[str, Any] = {}
    if "District Summary" in wb.sheetnames:
        s_sum = wb["District Summary"]
        for r in range(2, s_sum.max_row + 1):
            d_val = s_sum.cell(r, 2).value
            cnt_val = s_sum.cell(r, 3).value
            if d_val is not None and str(d_val).strip():
                d_clean = str(d_val).strip()
                try:
                    summary_counts[d_clean] = int(cnt_val) if cnt_val is not None else None
                except (ValueError, TypeError):
                    summary_counts[d_clean] = None

    mappings = []
    errors = []

    # Check for Block Mapping sheet
    if "Block Mapping" in wb.sheetnames:
        sheet = wb["Block Mapping"]
        for r in range(2, sheet.max_row + 1):
            d_val = sheet.cell(r, 2).value
            b_val = sheet.cell(r, 3).value
            if d_val is None and b_val is None:
                continue
            if not d_val or not str(d_val).strip():
                errors.append(f"Row {r}: District is missing.")
                continue
            if not b_val or not str(b_val).strip():
                errors.append(f"Row {r}: Block name is missing for District '{d_val}'.")
                continue
            mappings.append((str(d_val).strip(), str(b_val).strip()))
    else:
        # Fallback to active sheet
        sheet = wb.active
        h_row = [str(sheet.cell(1, c).value or "").strip().lower() for c in range(1, sheet.max_column + 1)]
        d_col = None
        b_col = None
        for idx, h in enumerate(h_row, 1):
            if "district" in h:
                d_col = idx
            elif "block" in h:
                b_col = idx
        if not d_col or not b_col:
            d_col, b_col = 1, 2

        for r in range(2, sheet.max_row + 1):
            d_val = sheet.cell(r, d_col).value
            b_val = sheet.cell(r, b_col).value
            if d_val is None and b_val is None:
                continue
            if not d_val or not str(d_val).strip():
                errors.append(f"Row {r}: District is missing.")
                continue
            if not b_val or not str(b_val).strip():
                errors.append(f"Row {r}: Block name is missing.")
                continue
            mappings.append((str(d_val).strip(), str(b_val).strip()))

    # Deduplicate within same district
    unique_mappings = []
    seen = set()
    for d, b in mappings:
        pair = (d, b)
        if pair not in seen:
            seen.add(pair)
            unique_mappings.append(pair)

    # Cross-validate against summary_counts if present
    d_blocks_map: Dict[str, list] = {}
    for d, b in unique_mappings:
        d_blocks_map.setdefault(d, []).append(b)

    for d, expected_cnt in summary_counts.items():
        if d not in d_blocks_map:
            errors.append(f"District '{d}' from Summary is missing in Block Mapping.")
        elif expected_cnt is not None and len(d_blocks_map[d]) != expected_cnt:
            errors.append(f"District '{d}' block count mismatch: Summary expects {expected_cnt}, found {len(d_blocks_map[d])}.")

    return {
        "state_name": default_state,
        "districts_count": len(d_blocks_map),
        "blocks_count": len(unique_mappings),
        "d_blocks_map": d_blocks_map,
        "errors": errors,
    }


def import_location_excel(
    file_bytes_or_path: Union[bytes, str],
    state_name: str = "Rajasthan",
    state_code: str = "RJ",
    dry_run: bool = False,
) -> Dict[str, Any]:
    """Import Location Master data from Excel with idempotency and cross-validation."""
    parsed = parse_location_workbook(file_bytes_or_path, default_state=state_name)
    d_blocks_map = parsed["d_blocks_map"]
    errors = list(parsed["errors"])

    total_districts = len(d_blocks_map)
    total_blocks = sum(len(blocks) for blocks in d_blocks_map.values())

    created_districts = 0
    skipped_districts = 0
    created_blocks = 0
    skipped_blocks = 0

    if dry_run:
        return {
            "state": state_name,
            "total_districts": total_districts,
            "total_blocks": total_blocks,
            "created_districts": total_districts,
            "skipped_districts": 0,
            "created_blocks": total_blocks,
            "skipped_blocks": 0,
            "invalid_records": len(errors),
            "errors": errors,
            "dry_run": True,
        }

    # Retrieve or create state
    state = db.session.execute(
        select(LocationState).where(func.lower(LocationState.name) == state_name.strip().lower())
    ).scalar_one_or_none()

    if not state:
        state = LocationState(
            name=state_name.strip(),
            code=state_code.strip() if state_code else None,
            is_active=True,
        )
        db.session.add(state)
        db.session.flush()

    # Pre-fetch existing districts for this state
    existing_districts = {
        d.name.lower(): d
        for d in db.session.execute(
            select(LocationDistrict).where(LocationDistrict.state_id == state.id)
        ).scalars().all()
    }

    # Process districts
    for dist_name, blocks in d_blocks_map.items():
        norm_dist = dist_name.strip().lower()
        district = existing_districts.get(norm_dist)
        if not district:
            district = LocationDistrict(
                state_id=state.id,
                name=dist_name.strip(),
                is_active=True,
            )
            db.session.add(district)
            db.session.flush()
            existing_districts[norm_dist] = district
            created_districts += 1
        else:
            skipped_districts += 1

        # Pre-fetch existing blocks for this district
        existing_blocks = {
            b.name.lower(): b
            for b in db.session.execute(
                select(LocationBlock).where(LocationBlock.district_id == district.id)
            ).scalars().all()
        }

        # Process blocks
        for blk_name in blocks:
            norm_blk = blk_name.strip().lower()
            if norm_blk not in existing_blocks:
                block = LocationBlock(
                    district_id=district.id,
                    name=blk_name.strip(),
                    is_active=True,
                )
                db.session.add(block)
                existing_blocks[norm_blk] = block
                created_blocks += 1
            else:
                skipped_blocks += 1

    db.session.commit()

    return {
        "state": state.name,
        "state_id": str(state.id),
        "total_districts": total_districts,
        "total_blocks": total_blocks,
        "created_districts": created_districts,
        "skipped_districts": skipped_districts,
        "created_blocks": created_blocks,
        "skipped_blocks": skipped_blocks,
        "invalid_records": len(errors),
        "errors": errors,
        "dry_run": False,
    }

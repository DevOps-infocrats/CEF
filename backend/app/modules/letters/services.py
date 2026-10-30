import os
import re
import uuid
import shutil
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from datetime import datetime, timezone
from decimal import Decimal
from typing import Dict, Any, Optional, List, Tuple
from flask import current_app
from sqlalchemy import select, func, update
from app.extensions import db
from app.common.database.transactions import atomic
from app.common.errors.exceptions import ValidationError, NotFoundError, ForbiddenError
from app.modules.auth.models import AuditLog, User
from app.modules.candidates.models import (
    CandidateApplication,
    APPLICATION_STATUS_APPROVED,
    APPLICATION_STATUS_OFFER_ACCEPTED,
)
from app.modules.letters.models import (
    DocumentTemplate,
    CandidateLetter,
    TEMPLATE_TYPE_LOI,
    TEMPLATE_TYPE_OFFER,
    VALID_TEMPLATE_TYPES,
    LETTER_STATUS_DRAFT,
    LETTER_STATUS_FINAL,
    LETTER_STATUS_SENT,
    LETTER_STATUS_ACCEPTED,
    LETTER_STATUS_FAILED,
    DEFAULT_LOI_MAPPINGS,
    DEFAULT_OFFER_MAPPINGS,
    utc_now,
)
import docx

try:
    from reportlab.lib.pagesizes import letter as LETTER_PAGESIZE
    from reportlab.pdfgen import canvas
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False


def replace_in_runs(paragraph, search_text: str, replacement_text: str) -> bool:
    """Safely replace text across runs in a paragraph while preserving font, size, bold, and styles."""
    if not search_text or search_text == replacement_text:
        return False

    replaced_any = False
    start_pos = 0

    while True:
        full_text = "".join(r.text for r in paragraph.runs)
        idx = full_text.find(search_text, start_pos)
        if idx == -1:
            break

        start_char = idx
        end_char = idx + len(search_text)

        cur = 0
        start_run = None
        start_offset = 0
        end_run = None
        end_offset = 0

        for i, r in enumerate(paragraph.runs):
            rlen = len(r.text)
            r_start = cur
            r_end = cur + rlen
            if start_run is None and r_start <= start_char < r_end:
                start_run = i
                start_offset = start_char - r_start
            if r_start < end_char <= r_end:
                end_run = i
                end_offset = end_char - r_start
                break
            cur = r_end

        if start_run is not None and end_run is not None:
            if start_run == end_run:
                r = paragraph.runs[start_run]
                r.text = r.text[:start_offset] + replacement_text + r.text[end_offset:]
            else:
                paragraph.runs[start_run].text = paragraph.runs[start_run].text[:start_offset] + replacement_text
                paragraph.runs[end_run].text = paragraph.runs[end_run].text[end_offset:]
                for mid in range(start_run + 1, end_run):
                    paragraph.runs[mid].text = ""
            replaced_any = True
            start_pos = start_char + len(replacement_text)
        else:
            break

    return replaced_any


def replace_in_cell(cell, search_text: str, replacement_text: str) -> bool:
    """Replace text in a table cell without clearing cell styling."""
    replaced = False
    for p in cell.paragraphs:
        if replace_in_runs(p, search_text, replacement_text):
            replaced = True
    return replaced


def update_cell_value_preserving_style(cell, new_value: str) -> None:
    """Set cell text while strictly preserving existing run formatting and XML properties."""
    new_str = str(new_value).strip()
    if not cell.paragraphs:
        cell.text = new_str
        return

    target_r_elem = None
    for p in cell.paragraphs:
        for r in p.runs:
            if r.text.strip():
                target_r_elem = r._r
                r.text = new_str
                break
        if target_r_elem is not None:
            break

    if target_r_elem is not None:
        for p in cell.paragraphs:
            for r in p.runs:
                if r._r is not target_r_elem:
                    r.text = ""
    else:
        if cell.paragraphs[0].runs:
            cell.paragraphs[0].runs[0].text = new_str
        else:
            cell.paragraphs[0].text = new_str


# Backward compatibility alias
set_cell_text_preserve_style = update_cell_value_preserving_style


def update_paragraph_preserving_style(paragraph, new_value: str) -> None:
    """Set paragraph text while preserving existing run formatting and XML properties."""
    new_str = str(new_value).strip()
    if not paragraph.runs:
        paragraph.text = new_str
        return
    paragraph.runs[0].text = new_str
    for r in paragraph.runs[1:]:
        r.text = ""


def apply_semantic_paragraph_edits(doc, edits: Dict[str, Any], letter_type: str) -> None:
    """
    Applies in-place semantic paragraph edits, note edits, and clause changes to candidate DOCX.
    Matches by semantic identifier, paragraph content pattern, or placeholder.
    """
    if not edits:
        return

    loi_matchers = {
        "header_tag": lambda txt: "letter of intent" in txt.lower(),
        "subject": lambda txt: "subject:" in txt.lower() or "offer letter" in txt.lower(),
        "salutation": lambda txt: txt.strip().startswith("Dear "),
        "p_intro": lambda txt: "pleased to inform you" in txt or "intends to engage your services" in txt,
        "p_location": lambda txt: "initial place of services" in txt,
        "p_start_date": lambda txt: "start date of services:" in txt.lower(),
        "p_comp_header": lambda txt: "compensation payable:" in txt.lower(),
        "p_perf_note": lambda txt: "performance linked component" in txt.lower(),
        "p_reserve_note": lambda txt: "reserves the right to revise" in txt.lower(),
        "p_transfer_h": lambda txt: "transferability:" in txt.lower(),
        "p_transfer_b": lambda txt: "services shall be transferable" in txt.lower(),
        "p_compliance_h": lambda txt: "compliance with company policies:" in txt.lower(),
        "p_compliance_b": lambda txt: "observe and comply with all applicable" in txt.lower(),
        "p_acceptance_h": lambda txt: "acceptance:" in txt.lower(),
        "p_acceptance_b1": lambda txt: "terms are acceptable to you" in txt.lower(),
        "p_acceptance_b2": lambda txt: "failure to report for duty" in txt.lower(),
        "p_acceptance_b3": lambda txt: "look forward to your association" in txt.lower(),
        "yours_faithfully": lambda txt: "yours faithfully" in txt.lower(),
        "for_company": lambda txt: "for fiber force" in txt.lower(),
        "authorized_signatory": lambda txt: "authorized signatory" in txt.lower(),
        "cand_acceptance_title": lambda txt: "acceptance by candidate" in txt.lower(),
        "cand_acceptance_text": lambda txt: "have read, understood" in txt.lower(),
        "cand_sig_label": lambda txt: txt.strip().startswith("Signature:"),
        "cand_name_label": lambda txt: txt.strip().startswith("Name:"),
        "cand_date_label": lambda txt: txt.strip().startswith("Date:"),
    }

    offer_matchers = {
        "header_tag": lambda txt: "appointment" in txt.lower() or "offer letter" in txt.lower(),
        "subject": lambda txt: "re: offer letter" in txt.lower() or (txt.strip().startswith("Re:") and "offer" in txt.lower()),
        "salutation": lambda txt: txt.strip().startswith("Dear "),
        "p10": lambda txt: "appointed as" in txt.lower(),
        "p11": lambda txt: "revisions in your salary" in txt.lower(),
        "p13": lambda txt: "appointment will be subject to the terms" in txt.lower(),
        "p14": lambda txt: "sign and return the duplicate copy" in txt.lower(),
        "p16": lambda txt: "join on or before" in txt.lower(),
        "p19": lambda txt: "look forward to working together" in txt.lower(),
        "yours_faithfully": lambda txt: "yours faithfully" in txt.lower(),
        "received_accepted": lambda txt: "received & accepted" in txt.lower(),
        "hr_dept": lambda txt: "hr department" in txt.lower(),
        "company_name": lambda txt: "fiber force private limited" in txt.lower(),
        "annex1_title": lambda txt: txt.strip() == "Annexure I" or ("annexure i" in txt.lower() and "ii" not in txt.lower()),
        "annex1_subtitle": lambda txt: "terms and conditions of appointment" in txt.lower(),
        "annex1_s1_h": lambda txt: "service subject to rules" in txt.lower(),
        "annex1_s1_b": lambda txt: "governed by the company" in txt.lower(),
        "annex1_s2_h": lambda txt: txt.strip().startswith("Probation"),
        "annex1_s3_h": lambda txt: "provident fund & esi" in txt.lower(),
        "annex1_s3_b": lambda txt: "provident fund as per" in txt.lower(),
        "annex1_s4_h": lambda txt: txt.strip() == "Transfer",
        "annex1_s4_b": lambda txt: "liable to be transferred" in txt.lower(),
        "annex1_s5_h": lambda txt: "transfer & mobility clause" in txt.lower(),
        "annex1_s5_b": lambda txt: "transferable to any office" in txt.lower(),
        "annex1_s6_h": lambda txt: "intellectual property rights" in txt.lower(),
        "annex1_s6_b": lambda txt: "invention, development" in txt.lower(),
        "annex1_s7_h": lambda txt: "moonlighting" in txt.lower(),
        "annex1_s7_b": lambda txt: "devote full time" in txt.lower() or "during the term of employment" in txt.lower(),
        "annex1_s8_h": lambda txt: "non-solicit" in txt.lower(),
        "annex1_s8_b": lambda txt: "12 months after cessation" in txt.lower(),
        "annex1_s9_h": lambda txt: "leaves and other service benefits" in txt.lower(),
        "annex1_s9_b": lambda txt: "entitled for leave" in txt.lower(),
        "annex1_s10_h": lambda txt: "certificate and testimonials" in txt.lower(),
        "annex1_s10_b": lambda txt: "certified copies of your certificates" in txt.lower(),
        "annex1_s11_h": lambda txt: txt.strip() == "Status",
        "annex1_s11_b": lambda txt: "informed about any change in your address" in txt.lower(),
        "annex1_s12_h": lambda txt: txt.strip().startswith("Secrecy"),
        "annex1_s13_h": lambda txt: txt.strip() == "Separation",
        "annex1_s13_sub1": lambda txt: "co-terminus employment" in txt.lower(),
        "annex2_title": lambda txt: "annexure" in txt.lower() and "ii" in txt.lower(),
        "annex2_note1": lambda txt: "income tax will be deducted" in txt.lower(),
        "annex2_note2": lambda txt: "prohibitory deductions" in txt.lower(),
        "annex2_note3": lambda txt: "mediclaim as per policy" in txt.lower(),
    }

    matchers = loi_matchers if letter_type == TEMPLATE_TYPE_LOI else offer_matchers

    for key, new_val in edits.items():
        if new_val is None or not str(new_val).strip():
            continue
        new_text = str(new_val).strip()

        # 1. Check if key is a semantic key
        matcher_fn = matchers.get(key)
        matched = False
        if matcher_fn:
            for p in doc.paragraphs:
                if p.text and matcher_fn(p.text):
                    update_paragraph_preserving_style(p, new_text)
                    matched = True
                    break

        if matched:
            continue

        # 2. Check direct paragraph placeholder {{key}}
        placeholder = f"{{{{{key}}}}}"
        placeholder_spaced = f"{{{{ {key} }}}}"
        for p in doc.paragraphs:
            if placeholder in p.text or placeholder_spaced in p.text:
                replace_in_runs(p, placeholder, new_text)
                replace_in_runs(p, placeholder_spaced, new_text)
                matched = True

        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    if placeholder in cell.text or placeholder_spaced in cell.text:
                        replace_in_cell(cell, placeholder, new_text)
                        replace_in_cell(cell, placeholder_spaced, new_text)
                        matched = True

        if matched:
            continue

        # 3. Direct substring search and replacement if key appears in text
        if len(key) >= 3:
            for p in doc.paragraphs:
                if key in p.text:
                    replace_in_runs(p, key, new_text)
            for t in doc.tables:
                for row in t.rows:
                    for cell in row.cells:
                        if key in cell.text:
                            replace_in_cell(cell, key, new_text)


def normalize_letter_layout_page_bounds(doc, letter_type: str) -> bool:
    """
    Ensure generated and edited candidate letters preserve their original intended page layout.
    For Letter of Intent (original 2-page document):
    1. Removes artificial consecutive blank paragraphs between Section 2 ('Transferability')
       and Section 3 ('Compliance with Company Policies:'). These blank paragraphs were placed
       by manual Enter key presses in the template; when dynamic candidate content (such as
       longer multi-line address, name, or location strings) expands on Page 1, these blank paragraphs
       spill onto Page 2, occupying 60-80pt of vertical space at the top of Page 2 and pushing
       the bottom acceptance block ('Date: ') onto an unintended 3rd page.
    2. Explicitly sets page_break_before on Section 3 ('3. Compliance with Company Policies:'),
       guaranteeing Section 3 starts cleanly at the top margin of Page 2 without wasted blank space.
    3. Strips any trailing empty paragraphs at the end of the document to avoid trailing blank pages.
    """
    changed = False
    is_loi = (letter_type == TEMPLATE_TYPE_LOI) or ("loi" in str(letter_type).lower())
    if is_loi:
        sec3_idx = None
        for idx, p in enumerate(doc.paragraphs):
            if "3. Compliance with Company Policies:" in p.text:
                sec3_idx = idx
                break

        if sec3_idx is not None:
            # Remove any empty paragraphs immediately preceding Section 3
            p_before_idx = sec3_idx - 1
            while p_before_idx >= 0 and not doc.paragraphs[p_before_idx].text.strip():
                empty_p = doc.paragraphs[p_before_idx]
                empty_p._p.getparent().remove(empty_p._p)
                p_before_idx -= 1
                changed = True

            # Ensure Section 3 starts cleanly on Page 2
            for p in doc.paragraphs:
                if "3. Compliance with Company Policies:" in p.text:
                    if not p.paragraph_format.page_break_before:
                        p.paragraph_format.page_break_before = True
                        changed = True
                    break

    # Strip any trailing blank paragraphs at document end
    while doc.paragraphs and not doc.paragraphs[-1].text.strip():
        last_p = doc.paragraphs[-1]
        last_p._p.getparent().remove(last_p._p)
        changed = True

    return changed


def prune_trailing_overflow_pages(pdf_path: str, max_expected_pages: int = 2) -> bool:
    """
    Ensure converted PDF does not exceed intended page bounds if conversion engine
    created an accidental trailing page containing only running header/footer or minimal text (< 200 chars).
    """
    if not pdf_path or not os.path.exists(pdf_path):
        return False
    try:
        import pypdf
        reader = pypdf.PdfReader(pdf_path)
        total_pages = len(reader.pages)
        if total_pages <= max_expected_pages:
            return False

        if total_pages == max_expected_pages + 1:
            last_page = reader.pages[-1]
            txt = last_page.extract_text().strip()
            is_spill = False
            if len(txt) < 200:
                is_spill = True
            elif "registered office" in txt.lower() and len(txt.replace("registered office", "").strip()) < 80:
                is_spill = True

            if is_spill:
                writer = pypdf.PdfWriter()
                for i in range(max_expected_pages):
                    writer.add_page(reader.pages[i])
                with open(pdf_path, "wb") as f:
                    writer.write(f)
                return True
    except Exception as e:
        if current_app:
            current_app.logger.warning(f"Could not prune trailing overflow page from PDF '{pdf_path}': {e}")
    return False


def convert_docx_to_pdf_libreoffice(docx_path: str, pdf_path: str) -> bool:
    """Convert a DOCX file to PDF using headless LibreOffice (Linux & Windows)."""
    import shutil
    import subprocess
    import tempfile

    candidate_bins = [
        shutil.which("libreoffice"),
        shutil.which("soffice"),
        "/usr/bin/libreoffice",
        "/usr/bin/soffice",
        "/usr/local/bin/libreoffice",
        "/usr/local/bin/soffice",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    lo_bin = None
    for cand in candidate_bins:
        if cand and os.path.exists(cand):
            lo_bin = cand
            break

    if not lo_bin:
        return False

    abs_docx = os.path.abspath(docx_path)
    abs_pdf = os.path.abspath(pdf_path)
    out_dir = os.path.dirname(abs_pdf)
    os.makedirs(out_dir, exist_ok=True)

    try:
        with tempfile.TemporaryDirectory() as temp_profile:
            cmd = [
                lo_bin,
                "--headless",
                f"-env:UserInstallation=file://{temp_profile.replace(os.sep, '/')}",
                "--convert-to",
                "pdf:writer_pdf_Export",
                abs_docx,
                "--outdir",
                out_dir,
            ]
            res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60)
            base_name = os.path.splitext(os.path.basename(abs_docx))[0]
            generated_pdf = os.path.join(out_dir, f"{base_name}.pdf")

            if os.path.exists(generated_pdf):
                if os.path.abspath(generated_pdf) != abs_pdf:
                    shutil.move(generated_pdf, abs_pdf)
                if os.path.exists(abs_pdf) and os.path.getsize(abs_pdf) > 0:
                    return True
            elif res.returncode != 0 and current_app:
                current_app.logger.warning(
                    f"LibreOffice conversion failed (code {res.returncode}): {res.stderr.decode('utf-8', errors='ignore')}"
                )
    except Exception as e:
        if current_app:
            current_app.logger.warning(f"LibreOffice conversion failed: {e}")

    return False


def convert_docx_to_pdf_native(docx_path: str, pdf_path: str, letter_type: Optional[str] = None) -> bool:
    """Convert a DOCX file to PDF preserving native fonts, tables, margins, and page breaks."""
    abs_docx = os.path.abspath(docx_path)
    abs_pdf = os.path.abspath(pdf_path)

    # 1. Primary for Windows: Word COM automation
    if os.name == "nt":
        try:
            import win32com.client
            import pythoncom

            pythoncom.CoInitialize()
            word = None
            doc = None
            try:
                word = win32com.client.DispatchEx("Word.Application")
                word.Visible = False
                word.DisplayAlerts = False
                doc = word.Documents.Open(abs_docx, ReadOnly=True)
                doc.ExportAsFixedFormat(abs_pdf, 17)  # 17 = wdExportFormatPDF
            finally:
                if doc:
                    try:
                        doc.Close(False)
                    except Exception:
                        pass
                if word:
                    try:
                        word.Quit()
                    except Exception:
                        pass
                pythoncom.CoUninitialize()

            if os.path.exists(abs_pdf) and os.path.getsize(abs_pdf) > 0:
                return True
        except Exception as e:
            if current_app:
                current_app.logger.warning(f"Word COM PDF export attempt failed: {e}")

    # 2. LibreOffice (Primary for Linux servers, robust secondary for Windows)
    try:
        if convert_docx_to_pdf_libreoffice(abs_docx, abs_pdf):
            return True
    except Exception as e:
        if current_app:
            current_app.logger.warning(f"LibreOffice conversion error: {e}")

    # 3. docx2pdf (if available on Windows/macOS)
    try:
        import docx2pdf

        docx2pdf.convert(abs_docx, abs_pdf)
        if os.path.exists(abs_pdf) and os.path.getsize(abs_pdf) > 0:
            return True
    except Exception as e:
        if current_app:
            current_app.logger.warning(f"docx2pdf conversion failed: {e}")

    # 4. Fallback: ReportLab
    if REPORTLAB_AVAILABLE:
        try:
            c = canvas.Canvas(abs_pdf, pagesize=LETTER_PAGESIZE)
            d = docx.Document(abs_docx)
            y = 750
            c.setFont("Helvetica-Bold", 12)
            c.drawString(50, y, "Official Candidate Document")
            y -= 25
            c.setFont("Helvetica", 9)
            for p in d.paragraphs:
                txt = p.text.strip()
                if txt:
                    for line in txt.split("\n"):
                        words = line.split()
                        cur_line = ""
                        for w in words:
                            if len(cur_line) + len(w) > 85:
                                c.drawString(50, y, cur_line)
                                y -= 13
                                cur_line = w + " "
                                if y < 45:
                                    c.showPage()
                                    y = 750
                                    c.setFont("Helvetica", 9)
                            else:
                                cur_line += w + " "
                        if cur_line:
                            c.drawString(50, y, cur_line)
                            y -= 13
                            if y < 45:
                                c.showPage()
                                y = 750
                                c.setFont("Helvetica", 9)
                    y -= 6
                    if y < 45:
                        c.showPage()
                        y = 750
                        c.setFont("Helvetica", 9)
            c.save()
            if os.path.exists(abs_pdf) and os.path.getsize(abs_pdf) > 0:
                pass
        except Exception as e:
            if current_app:
                current_app.logger.error(f"ReportLab fallback PDF failed: {e}")

    if os.path.exists(abs_pdf) and os.path.getsize(abs_pdf) > 0:
        if letter_type == TEMPLATE_TYPE_LOI or "loi" in os.path.basename(docx_path).lower():
            prune_trailing_overflow_pages(abs_pdf, max_expected_pages=2)
        return True

    return False


def resolve_storage_path(path: Optional[str]) -> str:
    """
    Cross-platform storage path resolver.
    Normalizes Windows laptop paths, Docker container paths, or relative paths
    to the active server's local storage directory.
    """
    if not path:
        return ""

    try:
        storage_root = os.path.abspath(os.path.join(current_app.root_path, "..", "storage"))
    except Exception:
        backend_storage = os.path.abspath(os.path.join(os.getcwd(), "backend", "storage"))
        if os.path.exists(backend_storage):
            storage_root = backend_storage
        else:
            storage_root = os.path.abspath(os.path.join(os.getcwd(), "storage"))

    # If the file already exists directly on disk as specified, return normalized path
    if os.path.exists(path):
        return os.path.abspath(path)

    # Clean separators
    p = str(path).replace("\\", "/")

    # Check for storage subfolders: generated_letters, templates, sent_emails, documents
    for sub in ("generated_letters", "templates", "sent_emails", "documents"):
        marker = f"/{sub}/"
        idx = p.rfind(marker)
        if idx != -1:
            rel = p[idx + 1 :]
            resolved = os.path.abspath(os.path.join(storage_root, *rel.split("/")))
            if os.path.exists(resolved):
                return resolved
            alt_backend = os.path.abspath(os.path.join(os.getcwd(), "backend", "storage", *rel.split("/")))
            if os.path.exists(alt_backend):
                return alt_backend
            os.makedirs(os.path.dirname(resolved), exist_ok=True)
            return resolved
        if p.startswith(f"{sub}/"):
            resolved = os.path.abspath(os.path.join(storage_root, *p.split("/")))
            if os.path.exists(resolved):
                return resolved
            alt_backend = os.path.abspath(os.path.join(os.getcwd(), "backend", "storage", *p.split("/")))
            if os.path.exists(alt_backend):
                return alt_backend
            os.makedirs(os.path.dirname(resolved), exist_ok=True)
            return resolved

    # Fallback to storage_root with basename
    fname = os.path.basename(p)
    resolved = os.path.abspath(os.path.join(storage_root, fname))
    if os.path.exists(resolved):
        return resolved
    alt_tmpl = os.path.abspath(os.path.join(storage_root, "templates", fname))
    if os.path.exists(alt_tmpl):
        return alt_tmpl
    alt_backend_tmpl = os.path.abspath(os.path.join(os.getcwd(), "backend", "storage", "templates", fname))
    if os.path.exists(alt_backend_tmpl):
        return alt_backend_tmpl
    os.makedirs(os.path.dirname(resolved), exist_ok=True)
    return resolved


def get_templates_storage_dir() -> str:
    try:
        path = os.path.join(current_app.root_path, "..", "storage", "templates")
    except Exception:
        path = os.path.join(os.getcwd(), "storage", "templates")
    os.makedirs(path, exist_ok=True)
    return os.path.abspath(path)


def get_generated_letters_storage_dir(app_id: str) -> str:
    try:
        path = os.path.join(current_app.root_path, "..", "storage", "generated_letters", str(app_id))
    except Exception:
        path = os.path.join(os.getcwd(), "storage", "generated_letters", str(app_id))
    os.makedirs(path, exist_ok=True)
    return os.path.abspath(path)


def get_sent_emails_storage_dir() -> str:
    try:
        path = os.path.join(current_app.root_path, "..", "storage", "sent_emails")
    except Exception:
        path = os.path.join(os.getcwd(), "storage", "sent_emails")
    os.makedirs(path, exist_ok=True)
    return os.path.abspath(path)


class EmailService:
    """Enterprise email delivery service with real SMTP support and isolated offline fallback."""

    @staticmethod
    def render_congratulations_email(
        candidate_name: str,
        application_number: str,
        registered_email: str,
        role_applied: str,
        attachments: List[Dict[str, Any]],
    ) -> tuple[str, str, str]:
        """Render the candidate congratulations email matching the approved HTML design.
        
        Returns:
            tuple[subject, html_body, text_body]
        """
        subject = "Congratulations on Your Selection | CEF Portal"

        # Generate HTML rows and text bullets for each document attachment
        doc_rows_html = []
        doc_bullets_text = []
        for att in attachments:
            fname = att.get("filename", "document.pdf")
            if "Letter_of_Intent" in fname:
                doc_title = "Letter of Intent"
            elif "Offer_Letter" in fname:
                doc_title = "Formal Offer Letter"
            else:
                doc_title = "Official Candidate Document"

            doc_bullets_text.append(f"  • {doc_title} ({fname})")

            doc_rows_html.append(f"""
                  <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="padding: 10px 0; border-bottom: 1px solid #e2e8f0;">
                    <tr>
                      <td width="36" valign="middle">
                        <div style="width: 28px; height: 32px; background-color: #fee2e2; border: 1px solid #fecaca; border-radius: 4px; text-align: center; line-height: 32px; font-size: 10px; font-weight: 800; color: #dc2626; font-family: Arial, sans-serif;">
                          PDF
                        </div>
                      </td>
                      <td valign="middle" style="padding-left: 10px;">
                        <div style="font-size: 13.5px; font-weight: 700; color: #0f172a;">
                          {doc_title}
                        </div>
                        <div style="font-size: 11.5px; color: #64748b; font-family: Consolas, monospace;">
                          {fname}
                        </div>
                      </td>
                      <td align="right" valign="middle">
                        <span style="font-size: 10.5px; font-weight: 600; color: #1e3a8a; background-color: #eff6ff; border: 1px solid #bfdbfe; padding: 3px 8px; border-radius: 4px;">
                          Official Copy
                        </span>
                      </td>
                    </tr>
                  </table>""")

        joined_doc_rows = "".join(doc_rows_html)
        doc_count = len(attachments)
        joined_doc_bullets = "\n".join(doc_bullets_text) if doc_bullets_text else "  • Finalized Document"

        html_body = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{subject}</title>
  <!--[if mso]>
  <style type="text/css">
    body, table, td, p, a, h1, h2, h3 {{ font-family: Arial, sans-serif !important; }}
    table {{ mso-table-lspace: 0pt; mso-table-rspace: 0pt; }}
    img {{ -ms-interpolation-mode: bicubic; }}
  </style>
  <![endif]-->
  <style type="text/css">
    body {{
      margin: 0;
      padding: 0;
      min-width: 100%;
      background-color: #f1f5f9;
      -webkit-text-size-adjust: 100%;
      -ms-text-size-adjust: 100%;
      font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
    }}
    table {{
      border-spacing: 0;
    }}
    td {{
      padding: 0;
    }}
    img {{
      border: 0;
    }}
    @media only screen and (max-width: 620px) {{
      .email-container {{
        width: 100% !important;
        max-width: 100% !important;
        border-radius: 0 !important;
        border-left: none !important;
        border-right: none !important;
      }}
      .header-cell {{
        padding: 32px 20px 28px 20px !important;
      }}
      .content-cell {{
        padding: 26px 20px !important;
      }}
      .heading-title {{
        font-size: 22px !important;
      }}
    }}
  </style>
</head>
<body style="margin: 0; padding: 0; background-color: #f1f5f9; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; color: #334155; line-height: 1.6;">

  <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="background-color: #f1f5f9; padding: 32px 12px 48px 12px; width: 100% !important;">
    <tr>
      <td align="center">

        <!-- Main Card Container -->
        <table role="presentation" class="email-container" width="600" cellpadding="0" cellspacing="0" border="0" style="width: 100%; max-width: 600px; background-color: #ffffff; border-radius: 14px; overflow: hidden; box-shadow: 0 10px 25px -5px rgba(15, 23, 42, 0.08), 0 8px 10px -6px rgba(15, 23, 42, 0.04); border: 1px solid #e2e8f0; margin: 0 auto;">

          <!-- 1. Premium Navy-Blue Gradient Congratulation Banner -->
          <tr>
            <td class="header-cell" style="background-color: #0b1329; background-image: linear-gradient(135deg, #0b1329 0%, #1e3a8a 55%, #0f172a 100%); padding: 42px 32px 36px 32px; text-align: center; border-bottom: 3px solid #f59e0b;">
              
              <!-- Elegant Gold Accent Pill with Congratulation Symbol -->
              <table role="presentation" cellpadding="0" cellspacing="0" border="0" align="center" style="margin: 0 auto 16px auto;">
                <tr>
                  <td style="display: inline-block; padding: 6px 16px; border-radius: 9999px; background-color: rgba(245, 158, 11, 0.16); border: 1px solid rgba(245, 158, 11, 0.45); color: #fbbf24; font-size: 12px; font-weight: 700; letter-spacing: 0.8px; text-transform: uppercase;">
                    <span style="color: #f59e0b; margin-right: 6px;">&#9733;</span> OFFICIAL SELECTION &amp; APPROVAL <span style="color: #f59e0b; margin-left: 6px;">&#9733;</span>
                  </td>
                </tr>
              </table>

              <!-- Main Heading -->
              <h1 class="heading-title" style="margin: 0; color: #ffffff; font-size: 25px; font-weight: 800; letter-spacing: -0.4px; line-height: 1.3;">
                Congratulations, {candidate_name}!
              </h1>

              <!-- Supporting Text -->
              <p style="margin: 10px 0 0 0; color: #cbd5e1; font-size: 14px; font-weight: 400; line-height: 1.5;">
                Your application has been approved by our HR team.
              </p>

              <!-- Application Meta Pill -->
              <table role="presentation" cellpadding="0" cellspacing="0" border="0" align="center" style="margin: 18px auto 0 auto;">
                <tr>
                  <td style="padding: 6px 14px; background-color: rgba(255, 255, 255, 0.08); border-radius: 8px; border: 1px solid rgba(255, 255, 255, 0.14); color: #cbd5e1; font-size: 11px; font-family: Consolas, 'Courier New', monospace;">
                    Application Ref: <strong style="color: #ffffff;">{application_number}</strong> &nbsp;|&nbsp; Role: <strong style="color: #ffffff;">{role_applied}</strong>
                  </td>
                </tr>
              </table>

            </td>
          </tr>

          <!-- 2. Clean Email Body -->
          <tr>
            <td class="content-cell" style="padding: 34px 32px 28px 32px; color: #334155; font-size: 14.5px; line-height: 1.65;">
              
              <!-- Salutation -->
              <p style="margin: 0 0 16px 0; font-size: 15.5px; font-weight: 700; color: #0f172a;">
                Dear {candidate_name},
              </p>

              <!-- Highlighted Selection Announcement -->
              <p style="margin: 0 0 16px 0; font-size: 16px; font-weight: 700; color: #1e3a8a;">
                Congratulations on your selection!
              </p>

              <p style="margin: 0 0 16px 0;">
                We are pleased to inform you that your application has been approved by our HR team. We congratulate you on this achievement and thank you for your interest in joining our organization.
              </p>

              <p style="margin: 0 0 22px 0;">
                Please find your finalized document(s) attached to this email for your reference.
              </p>

              <!-- 3. Clearly Structured 'Your Documents' Section -->
              <div style="margin: 0 0 24px 0;">
                <table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" style="margin-bottom: 8px;">
                  <tr>
                    <td>
                      <span style="font-size: 12px; font-weight: 800; text-transform: uppercase; letter-spacing: 0.8px; color: #475569;">
                        Your Documents
                      </span>
                    </td>
                    <td align="right">
                      <span style="font-size: 11px; font-weight: 600; color: #059669; background-color: #ecfdf5; border: 1px solid #a7f3d0; padding: 2px 8px; border-radius: 9999px;">
                        Finalized Attachments ({doc_count})
                      </span>
                    </td>
                  </tr>
                </table>

                <!-- Documents List Box -->
                <div style="background-color: #f8fafc; border: 1px solid #e2e8f0; border-radius: 10px; padding: 6px 14px;">
                  {joined_doc_rows}
                </div>
              </div>

              <!-- Important Instructions -->
              <p style="margin: 0 0 16px 0;">
                We request you to carefully review the attached documents, including the terms and conditions, role details, compensation, and any further instructions mentioned therein.
              </p>

              <p style="margin: 0 0 16px 0;">
                Your next steps, if any, will be communicated by our HR team. Should you require any clarification regarding your documents or the joining process, please feel free to contact us.
              </p>

              <p style="margin: 0 0 28px 0; font-weight: 500; color: #0f172a;">
                Once again, congratulations! We wish you success and look forward to the opportunity to work with you.
              </p>

              <!-- 4. Professional HR Signature -->
              <div style="border-top: 1px solid #e2e8f0; padding-top: 22px; margin-top: 8px;">
                <p style="margin: 0 0 6px 0; color: #64748b; font-size: 13.5px;">
                  Warm regards,
                </p>
                <p style="margin: 0 0 2px 0; font-size: 14.5px; font-weight: 800; color: #0f172a;">
                  Human Resources Team
                </p>
                <p style="margin: 0 0 2px 0; font-size: 13px; font-weight: 600; color: #1e3a8a;">
                  CEF Portal
                </p>
                <p style="margin: 0; font-size: 12.5px; color: #64748b;">
                  Candidate Recruitment &amp; HR Operations
                </p>
              </div>

            </td>
          </tr>

          <!-- 5. Subtle Footer -->
          <tr>
            <td style="background-color: #f8fafc; border-top: 1px solid #e2e8f0; padding: 18px 32px; text-align: center; font-size: 11.5px; color: #94a3b8; line-height: 1.5;">
              <p style="margin: 0 0 4px 0;">
                This official email was sent to <strong style="color: #64748b;">{registered_email}</strong>, which is registered with your application.
              </p>
              <p style="margin: 0; color: #cbd5e1;">
                Candidate Enrollment Form (CEF) Portal &middot; Confidential &amp; Proprietary
              </p>
            </td>
          </tr>

        </table>

      </td>
    </tr>
  </table>

</body>
</html>"""

        text_body = f"""Dear {candidate_name},

Congratulations on your selection!

We are pleased to inform you that your application has been approved by our HR team. We congratulate you on this achievement and thank you for your interest in joining our organization.

Please find your finalized document(s) attached to this email for your reference:
{joined_doc_bullets}

We request you to carefully review the attached documents, including the terms and conditions, role details, compensation, and any further instructions mentioned therein.

Your next steps, if any, will be communicated by our HR team. Should you require any clarification regarding your documents or the joining process, please feel free to contact us.

Once again, congratulations! We wish you success and look forward to the opportunity to work with you.

Warm regards,

Human Resources Team
CEF Portal
Candidate Recruitment & HR Operations

--------------------------------------------------------------------------------
This official email was sent to {registered_email}, which is registered with your application.
Application Ref: {application_number} | Role: {role_applied}
"""
        return subject, html_body, text_body

    @staticmethod
    def _resolve_smtp_config() -> Dict[str, Any]:
        """Resolve SMTP settings supporting standard environmental aliases and live .env updates."""
        try:
            from dotenv import load_dotenv
            possible_env_paths = [
                os.path.join(current_app.root_path, "..", ".env"),
                os.path.join(os.getcwd(), ".env"),
                os.path.join(os.getcwd(), "backend", ".env"),
            ]
            for ep in possible_env_paths:
                if os.path.exists(ep):
                    load_dotenv(dotenv_path=ep, override=True)
                    break
        except Exception:
            pass

        smtp_host = (
            os.getenv("SMTP_HOST")
            or os.getenv("MAIL_SERVER")
            or os.getenv("EMAIL_HOST")
            or ""
        ).strip()

        raw_port = (
            os.getenv("SMTP_PORT")
            or os.getenv("MAIL_PORT")
            or os.getenv("EMAIL_PORT")
            or "587"
        )
        try:
            smtp_port = int(raw_port)
        except ValueError:
            smtp_port = 587

        smtp_user = (
            os.getenv("SMTP_USERNAME")
            or os.getenv("SMTP_USER")
            or os.getenv("MAIL_USERNAME")
            or os.getenv("EMAIL_HOST_USER")
            or ""
        ).strip()

        smtp_pass = (
            os.getenv("SMTP_PASSWORD")
            or os.getenv("SMTP_PASS")
            or os.getenv("MAIL_PASSWORD")
            or os.getenv("EMAIL_HOST_PASSWORD")
            or ""
        ).strip()

        # If user pasted a Google App Password with spaces (e.g. 'xxxx xxxx xxxx xxxx'), clean them
        if "gmail" in smtp_host.lower() and " " in smtp_pass:
            smtp_pass = smtp_pass.replace(" ", "")

        smtp_use_tls = (
            os.getenv("SMTP_USE_TLS")
            or os.getenv("MAIL_USE_TLS")
            or os.getenv("EMAIL_USE_TLS")
            or "True"
        ).lower() in ("true", "1", "yes")

        smtp_use_ssl = (
            os.getenv("SMTP_USE_SSL")
            or os.getenv("MAIL_USE_SSL")
            or os.getenv("EMAIL_USE_SSL")
            or ("True" if smtp_port == 465 else "False")
        ).lower() in ("true", "1", "yes")

        default_sender = smtp_user if "@" in smtp_user else "hr@fiberforce.in"
        sender_email = (
            os.getenv("SMTP_SENDER_EMAIL")
            or os.getenv("MAIL_DEFAULT_SENDER")
            or os.getenv("DEFAULT_FROM_EMAIL")
            or os.getenv("EMAIL_FROM")
            or default_sender
        ).strip()

        sender_name = (
            os.getenv("SMTP_SENDER_NAME")
            or os.getenv("MAIL_SENDER_NAME")
            or "CEF Portal HR Operations"
        ).strip()

        return {
            "host": smtp_host,
            "port": smtp_port,
            "user": smtp_user,
            "password": smtp_pass,
            "use_tls": smtp_use_tls,
            "use_ssl": smtp_use_ssl,
            "sender_email": sender_email,
            "sender_name": sender_name,
        }

    @staticmethod
    def test_smtp_connection(to_test_email: Optional[str] = None) -> Dict[str, Any]:
        """Diagnose and verify live SMTP connection and optional test email dispatch."""
        cfg = EmailService._resolve_smtp_config()
        if not cfg["host"]:
            return {
                "success": False,
                "mode": "not_configured",
                "message": (
                    "SMTP is not configured in backend/.env. Currently running in local emulated mode. "
                    "Please set SMTP_HOST, SMTP_PORT, SMTP_USERNAME, and SMTP_PASSWORD to send live emails."
                ),
                "config": {
                    "host": None,
                    "port": cfg["port"],
                    "user": cfg["user"] or None,
                    "sender_email": cfg["sender_email"],
                },
            }

        import ssl
        context = ssl.create_default_context()
        smtp_verify_cert = os.getenv("SMTP_VERIFY_CERT", "True").lower() not in ("false", "0", "no")
        if not smtp_verify_cert:
            context.check_hostname = False
            context.verify_mode = ssl.CERT_NONE

        try:
            if cfg["use_ssl"] or cfg["port"] == 465:
                server = smtplib.SMTP_SSL(cfg["host"], cfg["port"], timeout=20, context=context)
                server.ehlo()
            else:
                server = smtplib.SMTP(cfg["host"], cfg["port"], timeout=20)
                server.ehlo()
                if cfg["use_tls"]:
                    server.starttls(context=context)
                    server.ehlo()

            if cfg["user"] and cfg["password"]:
                server.login(cfg["user"], cfg["password"])

            if to_test_email and "@" in to_test_email:
                msg = MIMEMultipart("alternative")
                msg["From"] = f"{cfg['sender_name']} <{cfg['sender_email']}>"
                msg["To"] = to_test_email
                msg["Subject"] = "SMTP Diagnostic Test | CEF Portal"
                msg["Date"] = datetime.now(timezone.utc).strftime("%a, %d %b %Y %H:%M:%S +0000")
                msg.attach(MIMEText("This is a live test email confirming that your CEF Portal SMTP configuration is active.", "plain", "utf-8"))
                server.send_message(msg)

            try:
                server.quit()
            except Exception:
                pass
            return {
                "success": True,
                "mode": "smtp",
                "message": f"Successfully connected and authenticated with SMTP server '{cfg['host']}:{cfg['port']}'.",
                "config": {
                    "host": cfg["host"],
                    "port": cfg["port"],
                    "user": cfg["user"],
                    "sender_email": cfg["sender_email"],
                    "tls": cfg["use_tls"],
                    "ssl": cfg["use_ssl"],
                },
            }
        except Exception as e:
            return {
                "success": False,
                "mode": "smtp_error",
                "message": f"SMTP Connection or Authentication Failed: {str(e)}",
                "config": {
                    "host": cfg["host"],
                    "port": cfg["port"],
                    "user": cfg["user"],
                    "sender_email": cfg["sender_email"],
                },
                "error": str(e),
            }

    @staticmethod
    def send_email(
        to_email: str,
        subject: str,
        html_body: str,
        text_body: str,
        attachments: Optional[List[Dict[str, Any]]] = None,
    ) -> Dict[str, Any]:
        """Send an email with attachments via SMTP or save to local storage in test/dev mode."""
        if not to_email or "@" not in to_email:
            raise ValidationError(f"Invalid recipient email address: '{to_email}'")

        cfg = EmailService._resolve_smtp_config()
        from_header = f"{cfg['sender_name']} <{cfg['sender_email']}>"

        msg = MIMEMultipart("mixed")
        msg["From"] = from_header
        msg["To"] = to_email
        msg["Subject"] = subject
        msg["Date"] = datetime.now(timezone.utc).strftime("%a, %d %b %Y %H:%M:%S +0000")

        # Body parts
        alt_part = MIMEMultipart("alternative")
        alt_part.attach(MIMEText(text_body, "plain", "utf-8"))
        alt_part.attach(MIMEText(html_body, "html", "utf-8"))
        msg.attach(alt_part)

        # Attachments
        if attachments:
            for att in attachments:
                fname = att.get("filename", "document.pdf")
                data = att.get("content")
                if data is None and "filepath" in att:
                    with open(att["filepath"], "rb") as f:
                        data = f.read()
                if data:
                    part = MIMEApplication(data, Name=fname)
                    part["Content-Disposition"] = f'attachment; filename="{fname}"'
                    msg.attach(part)

        # Emulated dispatch mode if SMTP is not configured or in test mode
        if not cfg["host"] or cfg["host"].lower() in ("console", "mock", "none") or current_app.config.get("TESTING"):
            sent_dir = get_sent_emails_storage_dir()
            msg_id = f"mock-msg-{uuid.uuid4().hex[:12]}"
            log_path = os.path.join(sent_dir, f"{msg_id}.eml")
            with open(log_path, "wb") as f:
                f.write(msg.as_bytes())

            current_app.logger.info(
                f"[EmailService:EMULATED] Dispatched email to {to_email} (Subject: '{subject}') -> Saved to {log_path}"
            )
            return {
                "success": True,
                "mode": "emulated",
                "message_id": msg_id,
                "recipient": to_email,
                "log_path": log_path,
                "smtp_host": None,
            }

        # Live SMTP delivery
        try:
            import ssl
            context = ssl.create_default_context()
            smtp_verify_cert = os.getenv("SMTP_VERIFY_CERT", "True").lower() not in ("false", "0", "no")
            if not smtp_verify_cert:
                context.check_hostname = False
                context.verify_mode = ssl.CERT_NONE

            if cfg["use_ssl"] or cfg["port"] == 465:
                server = smtplib.SMTP_SSL(cfg["host"], cfg["port"], timeout=30, context=context)
                server.ehlo()
            else:
                server = smtplib.SMTP(cfg["host"], cfg["port"], timeout=30)
                server.ehlo()
                if cfg["use_tls"]:
                    server.starttls(context=context)
                    server.ehlo()

            if cfg["user"] and cfg["password"]:
                server.login(cfg["user"], cfg["password"])

            server.send_message(msg)
            try:
                server.quit()
            except Exception:
                pass

            msg_id = f"smtp-msg-{uuid.uuid4().hex[:12]}"
            current_app.logger.info(f"[EmailService:SMTP] Dispatched email to {to_email} via {cfg['host']}:{cfg['port']}")
            return {
                "success": True,
                "mode": "smtp",
                "message_id": msg_id,
                "recipient": to_email,
                "smtp_host": cfg["host"],
            }
        except Exception as e:
            current_app.logger.error(f"[EmailService:ERROR] Failed to send email to {to_email}: {str(e)}")
            raise ValidationError(f"Email service delivery failure ({cfg['host']}:{cfg['port']}): {str(e)}")


class LetterTemplateService:
    """Manages master Word templates and their field placeholder mappings."""

    @staticmethod
    def ensure_default_templates() -> None:
        """Seed master LOI and Offer Letter v1 templates if not already present in the database, preserving all uploaded templates untouched."""
        storage_dir = get_templates_storage_dir()

        # Check Letter of Intent
        loi_count = db.session.execute(
            select(func.count(DocumentTemplate.id)).where(DocumentTemplate.template_type == TEMPLATE_TYPE_LOI)
        ).scalar() or 0

        loi_path = os.path.join(storage_dir, "letter_of_intent_v1.docx")
        if not os.path.exists(loi_path):
            doc = docx.Document()
            doc.add_heading("Letter of Intent", 0)
            doc.add_paragraph("Date: {{current_date}}")
            doc.add_paragraph("To,\n{{candidate_full_name}}\nAddress: {{candidate_address}}\nMobile: {{candidate_mobile}}\nEmail: {{candidate_email}}")
            doc.add_paragraph("Dear {{candidate_full_name}},\nWe intend to engage your services as {{approved_role_name}} in {{approved_location}}.")
            doc.add_paragraph("Start Date of Services: {{joining_date}}")
            doc.save(loi_path)

        if loi_count == 0:
            loi_tmpl = DocumentTemplate(
                template_type=TEMPLATE_TYPE_LOI,
                name="Letter of Intent Master Template",
                filename="letter_of_intent_v1.docx",
                storage_path=loi_path,
                version=1,
                is_active=True,
                field_mappings=DEFAULT_LOI_MAPPINGS,
            )
            db.session.add(loi_tmpl)

        # Check Offer Letter
        offer_count = db.session.execute(
            select(func.count(DocumentTemplate.id)).where(DocumentTemplate.template_type == TEMPLATE_TYPE_OFFER)
        ).scalar() or 0

        offer_path = os.path.join(storage_dir, "offer_letter_v1.docx")
        if not os.path.exists(offer_path):
            doc = docx.Document()
            doc.add_heading("Appointment / Offer Letter", 0)
            doc.add_paragraph("Ref: {{reference_number}}    Date: {{current_date}}")
            doc.add_paragraph("To,\nName: {{candidate_full_name}}\nMobile: {{candidate_mobile}}")
            doc.add_paragraph("Dear {{candidate_full_name}},\nYou have been appointed as {{approved_role_name}} in {{approved_location}}.")
            doc.add_paragraph("Joining Date: {{joining_date}}")
            doc.save(offer_path)

        if offer_count == 0:
            offer_tmpl = DocumentTemplate(
                template_type=TEMPLATE_TYPE_OFFER,
                name="Offer Letter Master Template",
                filename="offer_letter_v1.docx",
                storage_path=offer_path,
                version=1,
                is_active=True,
                field_mappings=DEFAULT_OFFER_MAPPINGS,
            )
            db.session.add(offer_tmpl)

        db.session.commit()

    @staticmethod
    def get_active_template(template_type: str) -> DocumentTemplate:
        if template_type not in VALID_TEMPLATE_TYPES:
            raise ValidationError(f"Invalid template type: '{template_type}'. Expected one of {list(VALID_TEMPLATE_TYPES)}")

        # Query the currently active template with the highest version
        tmpl = db.session.execute(
            select(DocumentTemplate)
            .where(DocumentTemplate.template_type == template_type)
            .where(DocumentTemplate.is_active.is_(True))
            .order_by(DocumentTemplate.version.desc(), DocumentTemplate.created_at.desc())
        ).scalars().first()

        if not tmpl:
            LetterTemplateService.ensure_default_templates()
            tmpl = db.session.execute(
                select(DocumentTemplate)
                .where(DocumentTemplate.template_type == template_type)
                .where(DocumentTemplate.is_active.is_(True))
                .order_by(DocumentTemplate.version.desc(), DocumentTemplate.created_at.desc())
            ).scalars().first()

        if not tmpl:
            raise NotFoundError(f"No active template found for '{template_type}'.")

        resolved = resolve_storage_path(tmpl.storage_path)
        if resolved and os.path.exists(resolved):
            if tmpl.storage_path != resolved:
                tmpl.storage_path = resolved
                db.session.commit()
            return tmpl

        # If resolved path does not exist directly, check by filename under storage_dir
        storage_dir = get_templates_storage_dir()
        if tmpl.storage_path:
            candidate_by_name = os.path.join(storage_dir, os.path.basename(tmpl.storage_path))
            if os.path.exists(candidate_by_name):
                tmpl.storage_path = candidate_by_name
                db.session.commit()
                return tmpl

        # Fallback to default v1 on disk without permanently overwriting database record
        base_filename = "letter_of_intent_v1.docx" if template_type == TEMPLATE_TYPE_LOI else "offer_letter_v1.docx"
        v1_path = os.path.join(storage_dir, base_filename)
        if os.path.exists(v1_path):
            current_app.logger.warning(
                f"Master template '{tmpl.filename}' (v{tmpl.version}) physical file not found at '{tmpl.storage_path}'. Falling back to default v1."
            )
            return tmpl

        return tmpl

    @staticmethod
    def upload_new_template(
        template_type: str,
        file,
        user: User,
    ) -> DocumentTemplate:
        if template_type not in VALID_TEMPLATE_TYPES:
            raise ValidationError(f"Invalid template type: '{template_type}'")

        if not file or not getattr(file, "filename", None):
            raise ValidationError("A valid master .docx file must be uploaded.")

        filename = file.filename
        if not filename.lower().endswith(".docx"):
            raise ValidationError("Master template must be a Microsoft Word (.docx) file.")

        storage_dir = get_templates_storage_dir()

        # Find maximum version across all existing templates of this type
        max_version = db.session.execute(
            select(func.max(DocumentTemplate.version))
            .where(DocumentTemplate.template_type == template_type)
        ).scalar() or 0

        new_version = max_version + 1
        safe_filename = f"{template_type}_v{new_version}_{uuid.uuid4().hex[:8]}.docx"
        target_path = os.path.join(storage_dir, safe_filename)

        # Save uploaded file
        file.save(target_path)

        # Validate that the file is a valid Word document
        try:
            doc = docx.Document(target_path)
        except Exception:
            if os.path.exists(target_path):
                os.remove(target_path)
            raise ValidationError("The uploaded file is corrupted or not a valid Microsoft Word (.docx) document.")

        # Discover placeholder mappings in the uploaded document
        import re
        discovered_placeholders = set()
        placeholder_regex = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")
        for p in doc.paragraphs:
            for match in placeholder_regex.findall(p.text):
                discovered_placeholders.add(match)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for match in placeholder_regex.findall(cell.text):
                        discovered_placeholders.add(match)

        default_map = (
            DEFAULT_LOI_MAPPINGS if template_type == TEMPLATE_TYPE_LOI else DEFAULT_OFFER_MAPPINGS
        )
        if discovered_placeholders:
            mappings = [
                {"placeholder": f"{{{{{k}}}}}", "label": k.replace("_", " ").title(), "description": f"Field '{k}'"}
                for k in sorted(discovered_placeholders)
            ]
        else:
            mappings = default_map

        # Atomically deactivate previous active templates and save new active master
        with atomic():
            all_prev = db.session.execute(
                select(DocumentTemplate)
                .where(DocumentTemplate.template_type == template_type)
                .where(DocumentTemplate.is_active.is_(True))
            ).scalars().all()
            for pt in all_prev:
                pt.is_active = False

            new_tmpl = DocumentTemplate(
                template_type=template_type,
                name=f"{'Letter of Intent' if template_type == TEMPLATE_TYPE_LOI else 'Offer Letter'} Master Template v{new_version}",
                filename=filename,
                storage_path=target_path,
                version=new_version,
                is_active=True,
                field_mappings=mappings,
                created_by_id=user.id,
            )
            db.session.add(new_tmpl)

            # Audit log
            audit = AuditLog(
                action="TEMPLATE_UPLOADED",
                user_id=user.id,
                resource="document_templates",
                resource_id=str(new_tmpl.id),
                details={
                    "template_type": template_type,
                    "version": new_version,
                    "filename": filename,
                },
            )
            db.session.add(audit)

        return new_tmpl


class LetterGenerationService:
    """Handles automated candidate data population, draft editing, docx/pdf generation, and email dispatch."""

    @staticmethod
    def calculate_salary_breakdown(monthly_salary: float) -> Dict[str, Any]:
        """Generate structured salary components for Annexure II."""
        if monthly_salary is None or monthly_salary <= 0:
            return {
                "monthly_ctc": "",
                "yearly_ctc": "",
                "total_ctc_monthly": "",
                "basic_monthly": "",
                "basic_yearly": "",
                "hra_monthly": "0",
                "hra_yearly": "0",
                "conveyance_monthly": "0",
                "conveyance_yearly": "0",
                "other_allowance_monthly": "",
                "other_allowance_yearly": "",
                "medical_monthly": "",
                "medical_yearly": "",
                "bonus_monthly": "",
                "bonus_yearly": "",
                "gratuity_monthly": "",
                "gratuity_yearly": "",
                "gross_monthly": "",
                "gross_yearly": "",
                "emp_pf_monthly": "0",
                "emp_pf_yearly": "0",
                "emp_esic_monthly": "0",
                "emp_esic_yearly": "0",
                "gpa_monthly": "",
                "gpa_yearly": "",
                "prof_tax_monthly": "0",
                "prof_tax_yearly": "0",
                "income_tax_monthly": "0",
                "income_tax_yearly": "0",
                "deductions_monthly": "",
                "deductions_yearly": "",
                "take_home_monthly": "",
                "take_home_yearly": "",
                "pli_monthly": "0",
                "pli_yearly": "0",
                "net_take_home_monthly": "",
                "net_take_home_yearly": "",
                "empr_pf_monthly": "0",
                "empr_pf_yearly": "0",
                "empr_esic_monthly": "0",
                "empr_esic_yearly": "0",
            }

        monthly_ctc = round(monthly_salary, 2)
        basic_monthly = round(monthly_ctc * 0.60, 2)
        medical_monthly = 45.00
        bonus_monthly = round(basic_monthly * 0.0833, 2)
        gratuity_monthly = round(basic_monthly * 0.0481, 2)

        other_allowance = round(monthly_ctc - basic_monthly - medical_monthly - bonus_monthly - gratuity_monthly, 2)
        if other_allowance < 0:
            other_allowance = 0.00

        gross_monthly = monthly_ctc
        gpa_monthly = 45.00
        deductions_monthly = gpa_monthly
        take_home_monthly = round(gross_monthly - deductions_monthly, 2)
        yearly_ctc = round(monthly_ctc * 12, 2)

        return {
            "monthly_ctc": f"{int(monthly_ctc)}",
            "yearly_ctc": f"{int(yearly_ctc)}",
            "total_ctc_monthly": f"{int(monthly_ctc)}",
            "basic_monthly": f"{int(basic_monthly)}",
            "basic_yearly": f"{int(basic_monthly * 12)}",
            "hra_monthly": "0",
            "hra_yearly": "0",
            "conveyance_monthly": "0",
            "conveyance_yearly": "0",
            "other_allowance_monthly": f"{int(other_allowance)}",
            "other_allowance_yearly": f"{int(other_allowance * 12)}",
            "medical_monthly": f"{int(medical_monthly)}",
            "medical_yearly": f"{int(medical_monthly * 12)}",
            "bonus_monthly": f"{int(bonus_monthly)}",
            "bonus_yearly": f"{int(bonus_monthly * 12)}",
            "gratuity_monthly": f"{int(gratuity_monthly)}",
            "gratuity_yearly": f"{int(gratuity_monthly * 12)}",
            "gross_monthly": f"{int(gross_monthly)}",
            "gross_yearly": f"{int(gross_monthly * 12)}",
            "emp_pf_monthly": "0",
            "emp_pf_yearly": "0",
            "emp_esic_monthly": "0",
            "emp_esic_yearly": "0",
            "gpa_monthly": f"{int(gpa_monthly)}",
            "gpa_yearly": f"{int(gpa_monthly * 12)}",
            "prof_tax_monthly": "0",
            "prof_tax_yearly": "0",
            "income_tax_monthly": "0",
            "income_tax_yearly": "0",
            "deductions_monthly": f"{int(deductions_monthly)}",
            "deductions_yearly": f"{int(deductions_monthly * 12)}",
            "take_home_monthly": f"{int(take_home_monthly)}",
            "take_home_yearly": f"{int(take_home_monthly * 12)}",
            "pli_monthly": "0",
            "pli_yearly": "0",
            "net_take_home_monthly": f"{int(take_home_monthly)}",
            "net_take_home_yearly": f"{int(take_home_monthly * 12)}",
            "empr_pf_monthly": "0",
            "empr_pf_yearly": "0",
            "empr_esic_monthly": "0",
            "empr_esic_yearly": "0",
        }

    @staticmethod
    def build_approved_location(app: CandidateApplication) -> str:
        """Format candidate deployment location as Circle, District, Block (e.g. 'Rajasthan, Ajmer, Block Arain')."""
        loc_parts = []
        circle = (app.circle or "").strip()
        if circle:
            loc_parts.append(circle)
        if app.district and app.district.strip():
            loc_parts.append(app.district.strip())
        if app.block and app.block.strip():
            b_val = app.block.strip()
            loc_parts.append(b_val if "block" in b_val.lower() else f"Block {b_val}")

        if not loc_parts and app.current_state:
            loc_parts.append(app.current_state.strip())

        return ", ".join(loc_parts) if loc_parts else (app.circle or app.current_state or "")

    @staticmethod
    def get_next_offer_reference_number(exclude_letter_id: Optional[uuid.UUID] = None) -> str:
        """Generate sequential offer letter reference number starting from 000100 (+1 for each release)."""
        now = datetime.now()
        all_offers = db.session.execute(
            select(CandidateLetter).where(CandidateLetter.letter_type == TEMPLATE_TYPE_OFFER)
        ).scalars().all()

        max_seq = 99  # Starts at 000100
        for l in all_offers:
            if exclude_letter_id and l.id == exclude_letter_id:
                continue
            if l.field_values and isinstance(l.field_values, dict):
                ref = str(l.field_values.get("reference_number", "")).strip()
                match = re.search(r"/(\d{6})$", ref)
                if match:
                    try:
                        num = int(match.group(1))
                        if num >= 100 and num > max_seq:
                            max_seq = num
                    except ValueError:
                        pass

        next_seq = max_seq + 1
        formatted_seq = f"{next_seq:06d}"
        return f"FFPL/RAJ/{now.strftime('%Y-%y')}/{now.strftime('%m')}/{formatted_seq}"

    @staticmethod
    def validate_and_normalize_joining_date(date_str: str) -> str:
        """Validate joining date format and normalize to standard DD/MM/YYYY."""
        clean = (date_str or "").strip()
        if not clean:
            raise ValidationError("Joining Date cannot be empty.")
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d"):
            try:
                dt = datetime.strptime(clean, fmt)
                return dt.strftime("%d/%m/%Y")
            except ValueError:
                continue
        raise ValidationError(f"Invalid date format: '{date_str}'. Expected valid calendar date in DD/MM/YYYY format.")

    @staticmethod
    def get_or_create_candidate_letters(
        application_id: uuid.UUID,
        user: User,
    ) -> Dict[str, Any]:
        """Retrieve existing letters for approved candidate or generate fresh drafts."""
        with atomic():
            app = db.session.execute(
                select(CandidateApplication)
                .where(CandidateApplication.id == application_id)
                .with_for_update()
            ).scalar_one_or_none()

            if not app:
                raise NotFoundError("Candidate application record not found.")

            if app.status != APPLICATION_STATUS_APPROVED:
                raise ValidationError(
                    f"Candidate application '{app.application_number}' has status '{app.status}'. "
                    "Letters can only be generated for APPROVED candidates."
                )

            # Query existing LOI and Offer Letter in one query
            existing_letters = db.session.execute(
                select(CandidateLetter)
                .where(CandidateLetter.application_id == app.id)
                .where(CandidateLetter.letter_type.in_([TEMPLATE_TYPE_LOI, TEMPLATE_TYPE_OFFER]))
                .order_by(CandidateLetter.version.desc(), CandidateLetter.created_at.desc())
            ).scalars().all()

            all_loi = [l for l in existing_letters if l.letter_type == TEMPLATE_TYPE_LOI]
            all_offer = [l for l in existing_letters if l.letter_type == TEMPLATE_TYPE_OFFER]

            loi_letter = all_loi[0] if all_loi else None
            offer_letter = all_offer[0] if all_offer else None

            # Clean up any unexpected duplicate records
            if len(all_loi) > 1:
                for dup in all_loi[1:]:
                    db.session.delete(dup)
            if len(all_offer) > 1:
                for dup in all_offer[1:]:
                    db.session.delete(dup)

            app_dir = get_generated_letters_storage_dir(str(app.id))
            resolved_loi = resolve_storage_path(loi_letter.docx_storage_path) if (loi_letter and loi_letter.docx_storage_path) else ""
            resolved_offer = resolve_storage_path(offer_letter.docx_storage_path) if (offer_letter and offer_letter.docx_storage_path) else ""

            needs_loi = not loi_letter or not resolved_loi or not os.path.exists(resolved_loi)
            needs_offer = not offer_letter or not resolved_offer or not os.path.exists(resolved_offer)

            if loi_letter and resolved_loi and loi_letter.docx_storage_path != resolved_loi:
                loi_letter.docx_storage_path = resolved_loi
            if offer_letter and resolved_offer and offer_letter.docx_storage_path != resolved_offer:
                offer_letter.docx_storage_path = resolved_offer

            if needs_loi or needs_offer:
                loi_tmpl = LetterTemplateService.get_active_template(TEMPLATE_TYPE_LOI)
                offer_tmpl = LetterTemplateService.get_active_template(TEMPLATE_TYPE_OFFER)

                sal_val = app.negotiated_salary if app.negotiated_salary is not None else (app.asking_salary if app.asking_salary is not None else app.current_salary)
                if sal_val is not None and float(sal_val) > 0:
                    monthly_salary_num = float(sal_val)
                    monthly_salary_str = f"{int(monthly_salary_num)}"
                    total_salary_str = f"{int(monthly_salary_num)}"
                    salary_breakdown = LetterGenerationService.calculate_salary_breakdown(monthly_salary_num)
                else:
                    monthly_salary_str = ""
                    total_salary_str = ""
                    salary_breakdown = LetterGenerationService.calculate_salary_breakdown(0)

                # Build candidate address safely without producing "None, None - None"
                addr_parts = []
                if app.current_address_line1 and app.current_address_line1.strip():
                    addr_parts.append(app.current_address_line1.strip())
                if app.current_address_line2 and app.current_address_line2.strip():
                    addr_parts.append(app.current_address_line2.strip())
                city_parts = [p.strip() for p in (app.current_city, app.current_state) if p and p.strip()]
                city_state = ", ".join(city_parts)
                if app.current_pincode and app.current_pincode.strip():
                    city_state = f"{city_state} - {app.current_pincode.strip()}" if city_state else app.current_pincode.strip()
                if city_state:
                    addr_parts.append(city_state)
                candidate_address = ", ".join(addr_parts)

                # Build approved deployment location (Circle, District, Block)
                approved_location = LetterGenerationService.build_approved_location(app)

                # Sequential Reference number (starts at 000100, increments +1 for each offer letter)
                ref_num = LetterGenerationService.get_next_offer_reference_number(
                    exclude_letter_id=offer_letter.id if offer_letter else None
                )

                default_fields = {
                    "current_date": datetime.now().strftime("%d/%m/%Y"),
                    "date": datetime.now().strftime("%d/%m/%Y"),
                    "candidate_full_name": app.full_name or "",
                    "candidate_name": app.full_name or "",
                    "candidate_address": candidate_address or "",
                    "candidate_mobile": app.mobile_number or "",
                    "mobile_number": app.mobile_number or "",
                    "candidate_email": app.email or "",
                    "candidate_aadhaar": app.aadhaar_number or "-",
                    "approved_role_name": app.role_applied or "",
                    "role_applied": app.role_applied or "",
                    "approved_location": approved_location or "",
                    "location": approved_location or "",
                    "joining_date": getattr(app, "joining_date", None) or (loi_letter.field_values.get("joining_date") if loi_letter and loi_letter.field_values else None) or (offer_letter.field_values.get("joining_date") if offer_letter and offer_letter.field_values else None) or "",
                    "reference_number": ref_num,
                    "monthly_salary": monthly_salary_str,
                    "performance_incentive": "0" if monthly_salary_str else "",
                    "miscellaneous_allowance": "0" if monthly_salary_str else "",
                    "total_salary": total_salary_str,
                    **salary_breakdown,
                }

                # Create LOI if missing
                if needs_loi:
                    loi_docx_path = os.path.join(app_dir, f"loi_{app.application_number}.docx")
                    LetterGenerationService._render_docx(loi_tmpl.storage_path, loi_docx_path, default_fields, TEMPLATE_TYPE_LOI)
                    if not loi_letter:
                        loi_letter = CandidateLetter(
                            application_id=app.id,
                            template_id=loi_tmpl.id,
                            letter_type=TEMPLATE_TYPE_LOI,
                            version=1,
                            status=LETTER_STATUS_DRAFT,
                            docx_storage_path=loi_docx_path,
                            field_values=default_fields,
                            last_edited_by_id=user.id,
                            last_edited_at=utc_now(),
                        )
                        db.session.add(loi_letter)
                    else:
                        loi_letter.docx_storage_path = loi_docx_path
                        loi_letter.field_values = default_fields
                    db.session.flush()

                # Create Offer Letter if missing
                if needs_offer:
                    offer_docx_path = os.path.join(app_dir, f"offer_{app.application_number}.docx")
                    LetterGenerationService._render_docx(offer_tmpl.storage_path, offer_docx_path, default_fields, TEMPLATE_TYPE_OFFER)
                    if not offer_letter:
                        offer_letter = CandidateLetter(
                            application_id=app.id,
                            template_id=offer_tmpl.id,
                            letter_type=TEMPLATE_TYPE_OFFER,
                            version=1,
                            status=LETTER_STATUS_DRAFT,
                            docx_storage_path=offer_docx_path,
                            field_values=default_fields,
                            last_edited_by_id=user.id,
                            last_edited_at=utc_now(),
                        )
                        db.session.add(offer_letter)
                    else:
                        offer_letter.docx_storage_path = offer_docx_path
                        offer_letter.field_values = default_fields
                    db.session.flush()

                    audit = AuditLog(
                        action="CANDIDATE_LETTERS_GENERATED",
                        user_id=user.id,
                        resource="candidate_applications",
                        resource_id=str(app.id),
                        details={
                            "application_number": app.application_number,
                            "candidate_id": str(app.id),
                        },
                    )
                    db.session.add(audit)

                db.session.commit()

        return {
            "application": {
                "id": str(app.id),
                "user_id": str(app.user_id),
                "application_number": app.application_number,
                "full_name": app.full_name,
                "email": app.email,
                "mobile_number": app.mobile_number,
                "role_applied": app.role_applied,
                "status": app.status,
                "district": app.district,
                "circle": app.circle,
                "current_state": app.current_state,
                "joining_date": getattr(app, "joining_date", None) or (loi_letter.field_values.get("joining_date") if loi_letter and loi_letter.field_values else None) or (offer_letter.field_values.get("joining_date") if offer_letter and offer_letter.field_values else None) or None,
                "negotiated_salary": float(app.negotiated_salary) if app.negotiated_salary else None,
                "offer_accepted_at": app.offer_accepted_at.isoformat() if app.offer_accepted_at else None,
                "offer_acceptance_status": app.offer_acceptance_status,
                "offer_accepted_by_id": str(app.offer_accepted_by_id) if app.offer_accepted_by_id else None,
            },
            "letter_of_intent": LetterGenerationService._serialize_letter(loi_letter),
            "offer_letter": LetterGenerationService._serialize_letter(offer_letter),
        }

    @staticmethod
    def ensure_candidate_docx_exists(letter: CandidateLetter) -> str:
        """Ensure candidate DOCX file exists on disk, regenerating from master template and fields if missing."""
        resolved_docx = resolve_storage_path(letter.docx_storage_path)
        if not resolved_docx or not os.path.exists(resolved_docx):
            app = letter.application or db.session.execute(
                select(CandidateApplication).where(CandidateApplication.id == letter.application_id)
            ).scalar_one_or_none()

            app_id_str = str(letter.application_id)
            app_dir = get_generated_letters_storage_dir(app_id_str)
            app_num = app.application_number if app else app_id_str[:8]
            prefix = "loi" if letter.letter_type == TEMPLATE_TYPE_LOI else "offer"
            resolved_docx = os.path.join(app_dir, f"{prefix}_{app_num}.docx")

            tmpl = letter.template or LetterTemplateService.get_active_template(letter.letter_type)
            tmpl_path = resolve_storage_path(tmpl.storage_path)
            LetterGenerationService._render_docx(
                tmpl_path,
                resolved_docx,
                letter.field_values or {},
                letter.letter_type,
            )
            # Re-apply any saved content_edits on top of the freshly rendered DOCX
            content_edits = (letter.field_values or {}).get("content_edits")
            if content_edits and os.path.exists(resolved_docx):
                try:
                    d = docx.Document(resolved_docx)
                    apply_semantic_paragraph_edits(d, content_edits, letter.letter_type)
                    d.save(resolved_docx)
                except Exception as e:
                    current_app.logger.warning(f"Could not re-apply content edits on regenerated DOCX: {e}")

            letter.docx_storage_path = resolved_docx
            db.session.commit()
        else:
            if letter.docx_storage_path != resolved_docx:
                letter.docx_storage_path = resolved_docx
                db.session.commit()
            if resolved_docx and os.path.exists(resolved_docx):
                try:
                    d = docx.Document(resolved_docx)
                    if normalize_letter_layout_page_bounds(d, letter.letter_type):
                        d.save(resolved_docx)
                except Exception as e:
                    current_app.logger.warning(f"Could not normalize layout bounds on existing DOCX: {e}")

        return resolved_docx

    @staticmethod
    def _serialize_letter(letter: CandidateLetter) -> Dict[str, Any]:
        page_count = 2 if letter.letter_type == TEMPLATE_TYPE_LOI else 5
        pdf_path = resolve_storage_path(letter.pdf_storage_path) if letter.pdf_storage_path else ""
        if pdf_path and os.path.exists(pdf_path):
            try:
                import pypdf
                reader = pypdf.PdfReader(pdf_path)
                page_count = len(reader.pages)
            except Exception:
                pass

        return {
            "id": str(letter.id),
            "letter_type": letter.letter_type,
            "version": letter.version,
            "status": letter.status,
            "page_count": page_count,
            "docx_url": f"/api/v1/applications/{letter.application_id}/letters/{letter.letter_type}/download?format=docx",
            "pdf_url": f"/api/v1/applications/{letter.application_id}/letters/{letter.letter_type}/download?format=pdf",
            "inline_pdf_url": f"/api/v1/applications/{letter.application_id}/letters/{letter.letter_type}/download?format=pdf&inline=1",
            "field_values": letter.field_values or {},
            "content": LetterGenerationService.extract_letter_content(letter),
            "last_edited_at": letter.last_edited_at.isoformat() if letter.last_edited_at else None,
            "sent_at": letter.sent_at.isoformat() if letter.sent_at else None,
            "accepted_at": letter.accepted_at.isoformat() if letter.accepted_at else None,
            "accepted_by_id": str(letter.accepted_by_id) if letter.accepted_by_id else None,
            "email_recipient": letter.email_recipient,
            "email_status": letter.email_status,
            "email_error": letter.email_error,
        }

    @staticmethod
    def _render_docx(template_path: str, output_path: str, field_values: Dict[str, Any], letter_type: Optional[str] = None) -> None:
        """Render Word document by copying master template and populating fields while strictly preserving all run styling."""
        resolved_template = resolve_storage_path(template_path)
        resolved_output = resolve_storage_path(output_path)

        is_loi = (letter_type == TEMPLATE_TYPE_LOI) or ("loi" in os.path.basename(resolved_output).lower())
        is_offer = (letter_type == TEMPLATE_TYPE_OFFER) or ("offer" in os.path.basename(resolved_output).lower())

        if not os.path.exists(resolved_template):
            storage_dir = get_templates_storage_dir()
            fallback_filename = "letter_of_intent_v1.docx" if is_loi else "offer_letter_v1.docx"
            fallback_path = os.path.join(storage_dir, fallback_filename)
            if os.path.exists(fallback_path):
                resolved_template = fallback_path
            else:
                LetterTemplateService.ensure_default_templates()
                resolved_template = fallback_path

        os.makedirs(os.path.dirname(os.path.abspath(resolved_output)), exist_ok=True)

        # Copy pristine master template to working copy path
        shutil.copy2(resolved_template, resolved_output)

        doc = docx.Document(resolved_output)

        cand_name = str(field_values.get("candidate_full_name", "")).strip()
        cand_mobile = str(field_values.get("candidate_mobile", "")).strip()
        cand_email = str(field_values.get("candidate_email", "")).strip()
        cand_aadhaar = str(field_values.get("candidate_aadhaar", "")).strip()
        cand_addr = str(field_values.get("candidate_address", "")).strip()
        approved_role = str(field_values.get("approved_role_name", "")).strip()
        approved_loc = str(field_values.get("approved_location", "")).strip()
        joining_date = str(field_values.get("joining_date", "")).strip()
        cur_date = str(field_values.get("current_date", "")).strip()
        ref_num = str(field_values.get("reference_number", "")).strip()
        monthly_sal = str(field_values.get("monthly_salary", "")).strip()

        # 1. Substitute any explicit {{placeholder}} variables across all document paragraphs
        for p in doc.paragraphs:
            for k, v in field_values.items():
                if v is not None and not isinstance(v, (dict, list)):
                    s_val = str(v)
                    replace_in_runs(p, f"{{{{{k}}}}}", s_val)
                    replace_in_runs(p, f"{{{{ {k} }}}}", s_val)

        # 2. Substitute any explicit {{placeholder}} variables across all table cells
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for k, v in field_values.items():
                        if v is not None and not isinstance(v, (dict, list)):
                            s_val = str(v)
                            replace_in_cell(cell, f"{{{{{k}}}}}", s_val)
                            replace_in_cell(cell, f"{{{{ {k} }}}}", s_val)

        if is_loi:
            # 1. Letter of Intent Paragraphs
            for p in doc.paragraphs:
                if cur_date:
                    replace_in_runs(p, "Date: 23/10/2026", f"Date: {cur_date}")
                    replace_in_runs(p, "23/10/2026", cur_date)
                if cand_name:
                    replace_in_runs(p, "Mr./Ms. Mangeet Singh", f"Mr./Ms. {cand_name}")
                    replace_in_runs(p, "Mangeet Singh", cand_name)
                    replace_in_runs(p, "Dear Mr. Mangeet Singh,", f"Dear {cand_name},")
                else:
                    replace_in_runs(p, "Mr./Ms. Mangeet Singh", "Mr./Ms. ")
                    replace_in_runs(p, "Mangeet Singh", "")
                    replace_in_runs(p, "Dear Mr. Mangeet Singh,", "Dear Candidate,")

                if cand_addr:
                    replace_in_runs(p, "Address: Bhatinda - Punjab", f"Address: {cand_addr}")
                    replace_in_runs(p, "Bhatinda - Punjab", cand_addr)
                else:
                    replace_in_runs(p, "Address: Bhatinda - Punjab", "Address: ")
                    replace_in_runs(p, "Bhatinda - Punjab", "")

                if cand_mobile:
                    replace_in_runs(p, "Mobile: 7737383505", f"Mobile: {cand_mobile}")
                    replace_in_runs(p, "7737383505", cand_mobile)
                else:
                    replace_in_runs(p, "Mobile: 7737383505", "Mobile: ")
                    replace_in_runs(p, "7737383505", "")

                if cand_email:
                    replace_in_runs(p, "Email ID: -", f"Email ID: {cand_email}")
                else:
                    replace_in_runs(p, "Email ID: -", "Email ID: ")

                if cand_aadhaar and cand_aadhaar != "-":
                    replace_in_runs(p, "Aadhar Number: -", f"Aadhar Number: {cand_aadhaar}")
                else:
                    replace_in_runs(p, "Aadhar Number: -", "Aadhar Number: ")

                if approved_role:
                    replace_in_runs(p, "HDD Operator", approved_role)
                else:
                    replace_in_runs(p, "HDD Operator", "")

                if approved_loc:
                    replace_in_runs(p, "District Sri Ganganagar State Rajasthan.", f"{approved_loc}.")
                    replace_in_runs(p, "District Sri Ganganagar State Rajasthan", approved_loc)
                else:
                    replace_in_runs(p, "District Sri Ganganagar State Rajasthan.", "")
                    replace_in_runs(p, "District Sri Ganganagar State Rajasthan", "")

                if joining_date:
                    replace_in_runs(p, "Start Date of Services: 01/11/2026", f"Start Date of Services: {joining_date}")
                    replace_in_runs(p, "01/11/2026", joining_date)
                else:
                    replace_in_runs(p, "Start Date of Services: 01/11/2026", "Start Date of Services: ")
                    replace_in_runs(p, "01/11/2026", "")

            # LOI Table (Salary & Total)
            for t in doc.tables:
                for row in t.rows:
                    for cell in row.cells:
                        replace_in_cell(cell, "40000", monthly_sal)

        elif is_offer:
            # 1. Offer Letter Paragraphs
            for p in doc.paragraphs:
                if ref_num:
                    replace_in_runs(p, "FFPL/RAJ/2026-27/10/00035", ref_num)
                if cur_date:
                    replace_in_runs(p, "Date: 20th Oct, 2026", f"Date: {cur_date}")
                    replace_in_runs(p, "20th Oct, 2026", cur_date)
                if cand_name:
                    replace_in_runs(p, "Name: Mr. Setu Jat", f"Name: {cand_name}")
                    replace_in_runs(p, "Name:  Mr. Setu Jat", f"Name: {cand_name}")
                    replace_in_runs(p, "Mr. Setu Jat", cand_name)
                    replace_in_runs(p, "Setu Jat", cand_name)
                    replace_in_runs(p, "Dear Mr. Setu Jat,", f"Dear {cand_name},")
                else:
                    replace_in_runs(p, "Name: Mr. Setu Jat", "Name: ")
                    replace_in_runs(p, "Name:  Mr. Setu Jat", "Name: ")
                    replace_in_runs(p, "Mr. Setu Jat", "")
                    replace_in_runs(p, "Setu Jat", "")
                    replace_in_runs(p, "Dear Mr. Setu Jat,", "Dear Candidate,")

                if cand_mobile:
                    replace_in_runs(p, "Mobile -+91 9571024426", f"Mobile: {cand_mobile}")
                    replace_in_runs(p, "9571024426", cand_mobile)
                else:
                    replace_in_runs(p, "Mobile -+91 9571024426", "Mobile: ")
                    replace_in_runs(p, "9571024426", "")

                if approved_role:
                    replace_in_runs(p, "HDD - Tracker", approved_role)
                else:
                    replace_in_runs(p, "HDD - Tracker", "")

                if approved_loc:
                    replace_in_runs(p, "Presently you will be posted in Behrunda 2 Rajasthan", f"Presently you will be posted in {approved_loc}")
                    replace_in_runs(p, "Behrunda 2 Rajasthan", approved_loc)
                else:
                    replace_in_runs(p, "Presently you will be posted in Behrunda 2 Rajasthan", "Presently you will be posted in -")
                    replace_in_runs(p, "Behrunda 2 Rajasthan", "")

                if joining_date:
                    replace_in_runs(p, "Dated: 01/10/2026.", f"Dated: {joining_date}.")
                    replace_in_runs(p, "01/10/2026", joining_date)
                else:
                    replace_in_runs(p, "Dated: 01/10/2026.", "Dated: ")
                    replace_in_runs(p, "01/10/2026", "")

            # Offer Letter Table 0 (Annexure II)
            if doc.tables:
                t = doc.tables[0]
                if cand_name and len(t.rows) > 0 and len(t.rows[0].cells) > 1:
                    replace_in_cell(t.rows[0].cells[1], "Mr. Setu Jat", cand_name)
                    replace_in_cell(t.rows[0].cells[1], "Setu Jat", cand_name)

                # Update structured Annexure II compensation rows while preserving cell formatting
                if len(t.rows) >= 28:
                    try:
                        ms_num = float(monthly_sal) if monthly_sal else 0.0
                        if ms_num > 0:
                            sb = LetterGenerationService.calculate_salary_breakdown(ms_num)
                            default_yearly = str(int(ms_num * 12))
                        else:
                            sb = {}
                            default_yearly = ""
                    except (ValueError, TypeError):
                        sb = {}
                        default_yearly = ""

                    update_cell_value_preserving_style(t.rows[3].cells[1], str(field_values.get("total_ctc_monthly", sb.get("total_ctc_monthly", monthly_sal))))
                    update_cell_value_preserving_style(t.rows[3].cells[2], str(field_values.get("yearly_ctc", sb.get("yearly_ctc", default_yearly))))
                    update_cell_value_preserving_style(t.rows[4].cells[1], str(field_values.get("total_ctc_monthly", sb.get("total_ctc_monthly", monthly_sal))))
                    update_cell_value_preserving_style(t.rows[4].cells[2], str(field_values.get("yearly_ctc", sb.get("yearly_ctc", default_yearly))))

                    for r_idx in range(8, min(len(t.rows), 28)):
                        row = t.rows[r_idx]
                        head_name = row.cells[0].text.strip().lower()
                        if "basic" in head_name:
                            update_cell_value_preserving_style(row.cells[1], str(field_values.get("basic_monthly", sb.get("basic_monthly", row.cells[1].text))))
                            update_cell_value_preserving_style(row.cells[2], str(field_values.get("basic_yearly", sb.get("basic_yearly", row.cells[2].text))))
                        elif "others allowances" in head_name:
                            update_cell_value_preserving_style(row.cells[1], str(field_values.get("other_allowance_monthly", sb.get("other_allowance_monthly", row.cells[1].text))))
                            update_cell_value_preserving_style(row.cells[2], str(field_values.get("other_allowance_yearly", sb.get("other_allowance_yearly", row.cells[2].text))))
                        elif "medical allowance" in head_name:
                            update_cell_value_preserving_style(row.cells[1], str(field_values.get("medical_monthly", sb.get("medical_monthly", row.cells[1].text))))
                            update_cell_value_preserving_style(row.cells[2], str(field_values.get("medical_yearly", sb.get("medical_yearly", row.cells[2].text))))
                        elif "advance bonus" in head_name:
                            update_cell_value_preserving_style(row.cells[1], str(field_values.get("bonus_monthly", sb.get("bonus_monthly", row.cells[1].text))))
                            update_cell_value_preserving_style(row.cells[2], str(field_values.get("bonus_yearly", sb.get("bonus_yearly", row.cells[2].text))))
                        elif "advance gratuity" in head_name:
                            update_cell_value_preserving_style(row.cells[1], str(field_values.get("gratuity_monthly", sb.get("gratuity_monthly", row.cells[1].text))))
                            update_cell_value_preserving_style(row.cells[2], str(field_values.get("gratuity_yearly", sb.get("gratuity_yearly", row.cells[2].text))))
                        elif "gross salary" in head_name:
                            update_cell_value_preserving_style(row.cells[1], str(field_values.get("gross_monthly", sb.get("gross_monthly", row.cells[1].text))))
                            update_cell_value_preserving_style(row.cells[2], str(field_values.get("gross_yearly", sb.get("gross_yearly", row.cells[2].text))))
                        elif "take home" in head_name and "net" not in head_name:
                            update_cell_value_preserving_style(row.cells[1], str(field_values.get("take_home_monthly", sb.get("take_home_monthly", row.cells[1].text))))
                            update_cell_value_preserving_style(row.cells[2], str(field_values.get("take_home_yearly", sb.get("take_home_yearly", row.cells[2].text))))
                        elif "net take home" in head_name:
                            update_cell_value_preserving_style(row.cells[1], str(field_values.get("net_take_home_monthly", sb.get("net_take_home_monthly", row.cells[1].text))))
                            update_cell_value_preserving_style(row.cells[2], str(field_values.get("net_take_home_yearly", sb.get("net_take_home_yearly", row.cells[2].text))))
                        elif "total ctc" in head_name:
                            update_cell_value_preserving_style(row.cells[1], str(field_values.get("total_ctc_monthly", sb.get("total_ctc_monthly", row.cells[1].text))))
                            update_cell_value_preserving_style(row.cells[2], str(field_values.get("yearly_ctc", sb.get("yearly_ctc", row.cells[2].text))))

        # Also support placeholder tags if custom template with {{placeholders}} was uploaded
        for k, v in field_values.items():
            if isinstance(v, (str, int, float)):
                target = f"{{{{{k}}}}}"
                rep = str(v) if v is not None else ""
                for p in doc.paragraphs:
                    replace_in_runs(p, target, rep)
                for t in doc.tables:
                    for row in t.rows:
                        for cell in row.cells:
                            replace_in_cell(cell, target, rep)

        # Ensure candidate document layout preserves page bounds (2 pages for LOI)
        normalize_letter_layout_page_bounds(doc, letter_type)
        doc.save(output_path)

    @staticmethod
    def extract_letter_content(letter: CandidateLetter) -> Dict[str, Any]:
        """Extract structured paragraphs, table cells, and fields from candidate DOCX."""
        docx_path = LetterGenerationService.ensure_candidate_docx_exists(letter)
        saved_content_edits = (letter.field_values or {}).get("content_edits", {})
        if not docx_path or not os.path.exists(docx_path):
            return {
                "paragraphs": {},
                "tables": [],
                "field_values": letter.field_values or {},
                "content_edits": saved_content_edits,
            }
        try:
            doc = docx.Document(docx_path)
            paragraphs = {str(i): p.text for i, p in enumerate(doc.paragraphs) if p.text.strip()}
            tables = [
                [[cell.text.strip() for cell in row.cells] for row in table.rows]
                for table in doc.tables
            ]
            return {
                "paragraphs": paragraphs,
                "tables": tables,
                "field_values": letter.field_values or {},
                "content_edits": saved_content_edits,
            }
        except Exception as e:
            current_app.logger.warning(f"Failed to extract DOCX content: {e}")
            return {
                "paragraphs": {},
                "tables": [],
                "field_values": letter.field_values or {},
                "content_edits": saved_content_edits,
            }

    @staticmethod
    def update_letter_draft(
        application_id: uuid.UUID,
        letter_type: str,
        field_values: Dict[str, Any],
        uploaded_docx: Optional[Any] = None,
        user: Optional[User] = None,
        text_edits: Optional[Dict[str, str]] = None,
        paragraphs: Optional[Dict[str, str]] = None,
        table_cells: Optional[Dict[str, str]] = None,
    ) -> CandidateLetter:
        """Update candidate letter draft with modified fields, in-place paragraph/cell edits, or revised DOCX upload."""
        letter = db.session.execute(
            select(CandidateLetter)
            .where(CandidateLetter.application_id == application_id)
            .where(CandidateLetter.letter_type == letter_type)
            .order_by(CandidateLetter.version.desc(), CandidateLetter.created_at.desc())
        ).scalars().first()

        if not letter:
            raise NotFoundError(f"Candidate letter for '{letter_type}' not found.")

        with atomic():
            docx_path = LetterGenerationService.ensure_candidate_docx_exists(letter)
            if uploaded_docx and uploaded_docx.filename:
                # Save revised docx directly
                uploaded_docx.save(docx_path)
                letter.docx_storage_path = docx_path
            else:
                doc = docx.Document(docx_path)
                old_fields = dict(letter.field_values or {})
                merged_fields = dict(old_fields)
                if field_values:
                    if "joining_date" in field_values:
                        raw_jd = field_values.get("joining_date")
                        if raw_jd is not None and str(raw_jd).strip():
                            normalized_jd = LetterGenerationService.validate_and_normalize_joining_date(str(raw_jd))
                            field_values["joining_date"] = normalized_jd
                        elif raw_jd == "":
                            field_values["joining_date"] = ""
                    merged_fields.update(field_values)

                # Persist content edits across sessions and reloads
                content_edits = dict(merged_fields.get("content_edits") or {})
                if text_edits:
                    content_edits.update(text_edits)
                if field_values and isinstance(field_values.get("content_edits"), dict):
                    content_edits.update(field_values["content_edits"])
                merged_fields["content_edits"] = content_edits

                # 1. If monthly_salary updated, recalculate salary breakdown
                if field_values and "monthly_salary" in field_values:
                    try:
                        ms = float(field_values["monthly_salary"])
                        sb = LetterGenerationService.calculate_salary_breakdown(ms)
                        merged_fields.update(sb)
                        merged_fields["total_salary"] = str(int(ms))
                        if letter_type == TEMPLATE_TYPE_LOI:
                            for t in doc.tables:
                                for row in t.rows:
                                    for cell in row.cells:
                                        replace_in_cell(cell, str(old_fields.get("monthly_salary", "40000")), str(int(ms)))
                        elif letter_type == TEMPLATE_TYPE_OFFER and doc.tables:
                            t = doc.tables[0]
                            if len(t.rows) >= 28:
                                update_cell_value_preserving_style(t.rows[3].cells[1], str(merged_fields.get("total_ctc_monthly", int(ms))))
                                update_cell_value_preserving_style(t.rows[3].cells[2], str(merged_fields.get("yearly_ctc", int(ms * 12))))
                                update_cell_value_preserving_style(t.rows[4].cells[1], str(merged_fields.get("total_ctc_monthly", int(ms))))
                                update_cell_value_preserving_style(t.rows[4].cells[2], str(merged_fields.get("yearly_ctc", int(ms * 12))))
                                for r_idx in range(8, min(len(t.rows), 28)):
                                    row = t.rows[r_idx]
                                    hname = row.cells[0].text.strip().lower()
                                    if "basic" in hname:
                                        update_cell_value_preserving_style(row.cells[1], str(merged_fields.get("basic_monthly", row.cells[1].text)))
                                        update_cell_value_preserving_style(row.cells[2], str(merged_fields.get("basic_yearly", row.cells[2].text)))
                                    elif "others allowances" in hname:
                                        update_cell_value_preserving_style(row.cells[1], str(merged_fields.get("other_allowance_monthly", row.cells[1].text)))
                                        update_cell_value_preserving_style(row.cells[2], str(merged_fields.get("other_allowance_yearly", row.cells[2].text)))
                                    elif "medical allowance" in hname:
                                        update_cell_value_preserving_style(row.cells[1], str(merged_fields.get("medical_monthly", "45")))
                                        update_cell_value_preserving_style(row.cells[2], str(merged_fields.get("medical_yearly", "540")))
                                    elif "advance bonus" in hname:
                                        update_cell_value_preserving_style(row.cells[1], str(merged_fields.get("bonus_monthly", row.cells[1].text)))
                                        update_cell_value_preserving_style(row.cells[2], str(merged_fields.get("bonus_yearly", row.cells[2].text)))
                                    elif "advance gratuity" in hname:
                                        update_cell_value_preserving_style(row.cells[1], str(merged_fields.get("gratuity_monthly", row.cells[1].text)))
                                        update_cell_value_preserving_style(row.cells[2], str(merged_fields.get("gratuity_yearly", row.cells[2].text)))
                                    elif "gross salary" in hname:
                                        update_cell_value_preserving_style(row.cells[1], str(merged_fields.get("gross_monthly", row.cells[1].text)))
                                        update_cell_value_preserving_style(row.cells[2], str(merged_fields.get("gross_yearly", row.cells[2].text)))
                                    elif "take home" in hname and "net" not in hname:
                                        update_cell_value_preserving_style(row.cells[1], str(merged_fields.get("take_home_monthly", row.cells[1].text)))
                                        update_cell_value_preserving_style(row.cells[2], str(merged_fields.get("take_home_yearly", row.cells[2].text)))
                                    elif "net take home" in hname:
                                        update_cell_value_preserving_style(row.cells[1], str(merged_fields.get("net_take_home_monthly", row.cells[1].text)))
                                        update_cell_value_preserving_style(row.cells[2], str(merged_fields.get("net_take_home_yearly", row.cells[2].text)))
                                    elif "total ctc" in hname:
                                        update_cell_value_preserving_style(row.cells[1], str(merged_fields.get("total_ctc_monthly", row.cells[1].text)))
                                        update_cell_value_preserving_style(row.cells[2], str(merged_fields.get("yearly_ctc", row.cells[2].text)))
                    except (ValueError, TypeError):
                        pass

                # 2. Update changed field values in paragraphs and tables
                if field_values:
                    for k, new_v in field_values.items():
                        old_v = old_fields.get(k)
                        if new_v is not None and str(new_v).strip() != str(old_v or "").strip():
                            s_new = str(new_v).strip()
                            s_old = str(old_v).strip() if old_v else ""
                            for p in doc.paragraphs:
                                replace_in_runs(p, f"{{{{{k}}}}}", s_new)
                                replace_in_runs(p, f"{{{{ {k} }}}}", s_new)
                                if s_old and len(s_old) >= 5 and s_old.lower() not in ("date", "name", "role", "city", "state", "user", "type"):
                                    replace_in_runs(p, s_old, s_new)
                                if k == "joining_date":
                                    if letter_type == TEMPLATE_TYPE_LOI:
                                        replace_in_runs(p, "Start Date of Services: 01/11/2026", f"Start Date of Services: {s_new}")
                                        replace_in_runs(p, "Start Date of Services: ", f"Start Date of Services: {s_new}")
                                        replace_in_runs(p, "Start Date of Services:", f"Start Date of Services: {s_new}")
                                        replace_in_runs(p, "01/11/2026", s_new)
                                        if s_old and len(s_old) >= 5:
                                            replace_in_runs(p, f"Start Date of Services: {s_old}", f"Start Date of Services: {s_new}")
                                            replace_in_runs(p, s_old, s_new)
                                    elif letter_type == TEMPLATE_TYPE_OFFER:
                                        replace_in_runs(p, "Dated: 01/10/2026.", f"Dated: {s_new}.")
                                        replace_in_runs(p, "Dated: .", f"Dated: {s_new}.")
                                        replace_in_runs(p, "Dated: ", f"Dated: {s_new}")
                                        replace_in_runs(p, "01/10/2026", s_new)
                                        if s_old and len(s_old) >= 5:
                                            replace_in_runs(p, f"Dated: {s_old}.", f"Dated: {s_new}.")
                                            replace_in_runs(p, s_old, s_new)
                            for t in doc.tables:
                                for row in t.rows:
                                    for cell in row.cells:
                                        replace_in_cell(cell, f"{{{{{k}}}}}", s_new)
                                        replace_in_cell(cell, f"{{{{ {k} }}}}", s_new)
                                        if s_old and len(s_old) >= 5 and s_old.lower() not in ("date", "name", "role", "city", "state", "user", "type"):
                                            replace_in_cell(cell, s_old, s_new)

                # 3. Direct paragraph index replacements
                if paragraphs:
                    for idx_str, new_text in paragraphs.items():
                        try:
                            idx = int(idx_str)
                            if 0 <= idx < len(doc.paragraphs):
                                p = doc.paragraphs[idx]
                                new_s = str(new_text)
                                if p.runs:
                                    p.runs[0].text = new_s
                                    for r in p.runs[1:]:
                                        r.text = ""
                                else:
                                    p.text = new_s
                        except (ValueError, IndexError):
                            pass

                # 4. Direct table cell replacements
                if table_cells:
                    for coord_str, new_text in table_cells.items():
                        try:
                            parts = [int(x) for x in coord_str.split(",")]
                            if len(parts) == 3:
                                ti, ri, ci = parts
                                if 0 <= ti < len(doc.tables) and 0 <= ri < len(doc.tables[ti].rows) and 0 <= ci < len(doc.tables[ti].rows[ri].cells):
                                    update_cell_value_preserving_style(doc.tables[ti].rows[ri].cells[ci], str(new_text))
                        except (ValueError, IndexError):
                            pass

                # 5. Semantic paragraph, note, and clause replacements
                if content_edits:
                    apply_semantic_paragraph_edits(doc, content_edits, letter_type)

                # 6. Direct text replacements
                if text_edits:
                    for old_t, new_t in text_edits.items():
                        if old_t and old_t != new_t and len(old_t) >= 3:
                            for p in doc.paragraphs:
                                replace_in_runs(p, old_t, str(new_t))
                            for t in doc.tables:
                                for row in t.rows:
                                    for cell in row.cells:
                                        replace_in_cell(cell, old_t, str(new_t))

                if letter_type == TEMPLATE_TYPE_LOI:
                    normalize_letter_layout_page_bounds(doc, letter_type)

                doc.save(docx_path)
                letter.docx_storage_path = docx_path
                letter.field_values = merged_fields

                # Synchronize joining_date to parent CandidateApplication and peer CandidateLetter
                if field_values and "joining_date" in field_values:
                    new_jd = merged_fields.get("joining_date")
                    if new_jd:
                        from app.modules.candidates.models import CandidateApplication
                        app_record = db.session.get(CandidateApplication, application_id)
                        if app_record:
                            app_record.joining_date = new_jd

                        peer_type = TEMPLATE_TYPE_OFFER if letter_type == TEMPLATE_TYPE_LOI else TEMPLATE_TYPE_LOI
                        peer_letter = db.session.execute(
                            select(CandidateLetter)
                            .where(CandidateLetter.application_id == application_id)
                            .where(CandidateLetter.letter_type == peer_type)
                            .order_by(CandidateLetter.version.desc())
                        ).scalars().first()

                        if peer_letter and peer_letter.id != letter.id:
                            peer_docx_path = LetterGenerationService.ensure_candidate_docx_exists(peer_letter)
                            peer_fields = dict(peer_letter.field_values or {})
                            old_peer_jd = peer_fields.get("joining_date", "")
                            if old_peer_jd != new_jd:
                                peer_fields["joining_date"] = new_jd
                                peer_letter.field_values = peer_fields
                                if os.path.exists(peer_docx_path):
                                    try:
                                        peer_doc = docx.Document(peer_docx_path)
                                        for p in peer_doc.paragraphs:
                                            replace_in_runs(p, "{{joining_date}}", new_jd)
                                            replace_in_runs(p, "{{ joining_date }}", new_jd)
                                            if old_peer_jd:
                                                replace_in_runs(p, old_peer_jd, new_jd)
                                            if peer_type == TEMPLATE_TYPE_LOI:
                                                replace_in_runs(p, "Start Date of Services: 01/11/2026", f"Start Date of Services: {new_jd}")
                                                replace_in_runs(p, "Start Date of Services: ", f"Start Date of Services: {new_jd}")
                                                replace_in_runs(p, "01/11/2026", new_jd)
                                                if old_peer_jd:
                                                    replace_in_runs(p, f"Start Date of Services: {old_peer_jd}", f"Start Date of Services: {new_jd}")
                                                    replace_in_runs(p, old_peer_jd, new_jd)
                                            else:
                                                replace_in_runs(p, "Dated: 01/10/2026.", f"Dated: {new_jd}.")
                                                replace_in_runs(p, "Dated: .", f"Dated: {new_jd}.")
                                                replace_in_runs(p, "Dated: ", f"Dated: {new_jd}")
                                                replace_in_runs(p, "01/10/2026", new_jd)
                                                if old_peer_jd:
                                                    replace_in_runs(p, f"Dated: {old_peer_jd}.", f"Dated: {new_jd}.")
                                                    replace_in_runs(p, old_peer_jd, new_jd)
                                        for t in peer_doc.tables:
                                            for row in t.rows:
                                                for cell in row.cells:
                                                    replace_in_cell(cell, "{{joining_date}}", new_jd)
                                                    replace_in_cell(cell, "{{ joining_date }}", new_jd)
                                                    if old_peer_jd:
                                                        replace_in_cell(cell, old_peer_jd, new_jd)
                                        if peer_type == TEMPLATE_TYPE_LOI:
                                            normalize_letter_layout_page_bounds(peer_doc, peer_type)
                                        peer_doc.save(peer_docx_path)
                                        peer_letter.docx_storage_path = peer_docx_path
                                    except Exception:
                                        pass
                                peer_pdf = resolve_storage_path(peer_letter.pdf_storage_path) if peer_letter.pdf_storage_path else ""
                                if peer_pdf and os.path.exists(peer_pdf):
                                    try:
                                        os.remove(peer_pdf)
                                    except Exception:
                                        pass
                                peer_letter.pdf_storage_path = None
                                peer_letter.last_edited_by_id = user.id if user else None
                                peer_letter.last_edited_at = utc_now()

            # Invalidate cached pdf to trigger re-generation on demand
            letter_pdf = resolve_storage_path(letter.pdf_storage_path) if letter.pdf_storage_path else ""
            if letter_pdf and os.path.exists(letter_pdf):
                try:
                    os.remove(letter_pdf)
                except Exception:
                    pass
            letter.pdf_storage_path = None
            letter.last_edited_by_id = user.id if user else None
            letter.last_edited_at = utc_now()
            letter.version += 1
            letter.status = LETTER_STATUS_DRAFT

        # Automatically re-render PDF so review page can immediately preview latest changes
        try:
            LetterGenerationService.get_or_render_pdf(letter)
        except Exception as e:
            current_app.logger.warning(f"Could not pre-render PDF for {letter_type}: {e}")

        return letter

    @staticmethod
    def get_or_render_pdf(letter: CandidateLetter) -> str:
        """Ensure a high-fidelity PDF version of the candidate letter exists and return its path."""
        docx_path = LetterGenerationService.ensure_candidate_docx_exists(letter)

        resolved_pdf = resolve_storage_path(letter.pdf_storage_path) if letter.pdf_storage_path else ""
        if resolved_pdf and os.path.exists(resolved_pdf):
            try:
                if os.path.getmtime(resolved_pdf) >= os.path.getmtime(docx_path):
                    # Ensure cached LOI PDF does not exceed 2 pages due to stale generation
                    if letter.letter_type == TEMPLATE_TYPE_LOI and PYPDF_AVAILABLE:
                        try:
                            reader = PdfReader(resolved_pdf)
                            if len(reader.pages) > 2:
                                current_app.logger.info(f"Invalidating stale {len(reader.pages)}-page cached LOI PDF: {resolved_pdf}")
                                resolved_pdf = ""
                        except Exception:
                            pass
                    if resolved_pdf:
                        letter.pdf_storage_path = resolved_pdf
                        return resolved_pdf
            except Exception:
                pass

        pdf_path = docx_path.replace(".docx", ".pdf")

        # In testing mode, use ReportLab to avoid COM processes in headless runners
        if current_app.config.get("TESTING") and REPORTLAB_AVAILABLE:
            try:
                c = canvas.Canvas(pdf_path, pagesize=LETTER_PAGESIZE)
                doc = docx.Document(docx_path)
                y = 750
                c.setFont("Helvetica-Bold", 14)
                c.drawString(50, y, f"{'Letter of Intent' if letter.letter_type == TEMPLATE_TYPE_LOI else 'Appointment / Offer Letter'}")
                y -= 30
                c.setFont("Helvetica", 10)
                for p in doc.paragraphs:
                    text = p.text.strip()
                    if text:
                        words = text.split()
                        line = ""
                        for word in words:
                            if len(line) + len(word) > 80:
                                c.drawString(50, y, line)
                                y -= 14
                                line = word + " "
                                if y < 50:
                                    c.showPage()
                                    y = 750
                                    c.setFont("Helvetica", 10)
                            else:
                                line += word + " "
                        if line:
                            c.drawString(50, y, line)
                            y -= 16
                            if y < 50:
                                c.showPage()
                                y = 750
                                c.setFont("Helvetica", 10)
                c.save()
                if os.path.exists(pdf_path):
                    if letter.letter_type == TEMPLATE_TYPE_LOI:
                        prune_trailing_overflow_pages(pdf_path, max_expected_pages=2)
                    letter.pdf_storage_path = pdf_path
                    db.session.commit()
                    return pdf_path
            except Exception as e:
                current_app.logger.error(f"Test PDF rendering error: {e}")

        # Native conversion (Windows Word COM, LibreOffice headless, docx2pdf, or ReportLab fallback)
        converted = convert_docx_to_pdf_native(docx_path, pdf_path, letter_type=letter.letter_type)
        if converted and os.path.exists(pdf_path):
            if letter.letter_type == TEMPLATE_TYPE_LOI:
                prune_trailing_overflow_pages(pdf_path, max_expected_pages=2)
            letter.pdf_storage_path = pdf_path
            db.session.commit()
            return pdf_path

        raise ValidationError("Failed to convert reviewed candidate letter to PDF. Please verify LibreOffice, Microsoft Word, or document converter.")

    @staticmethod
    def send_letters_by_email(
        application_id: uuid.UUID,
        user: User,
        force_resend: bool = False,
        document_scope: str = "both",
    ) -> Dict[str, Any]:
        """Validate, package, and email candidate's Letter of Intent and/or Offer Letter."""
        app = db.session.execute(
            select(CandidateApplication).where(CandidateApplication.id == application_id)
        ).scalar_one_or_none()

        if not app:
            raise NotFoundError("Application record not found.")

        if app.status != APPLICATION_STATUS_APPROVED:
            raise ValidationError(
                f"Candidate application '{app.application_number}' has status '{app.status}'. "
                "Letters can only be sent for APPROVED candidates."
            )

        if not app.email or "@" not in app.email:
            raise ValidationError(
                f"Cannot send letters via email because candidate '{app.application_number}' has no valid registered email address. "
                "Email dispatch is blocked until an email address is provided. The approved candidate status and generated letters remain preserved."
            )

        include_loi = document_scope in ("both", TEMPLATE_TYPE_LOI, "loi")
        include_offer = document_scope in ("both", TEMPLATE_TYPE_OFFER, "offer")

        loi_letter = None
        offer_letter = None

        if include_loi:
            loi_letter = db.session.execute(
                select(CandidateLetter)
                .where(CandidateLetter.application_id == app.id)
                .where(CandidateLetter.letter_type == TEMPLATE_TYPE_LOI)
                .order_by(CandidateLetter.version.desc(), CandidateLetter.created_at.desc())
            ).scalars().first()
            if not loi_letter:
                raise ValidationError("Letter of Intent draft has not been generated for review.")

        if include_offer:
            offer_letter = db.session.execute(
                select(CandidateLetter)
                .where(CandidateLetter.application_id == app.id)
                .where(CandidateLetter.letter_type == TEMPLATE_TYPE_OFFER)
                .order_by(CandidateLetter.version.desc(), CandidateLetter.created_at.desc())
            ).scalars().first()
            if not offer_letter:
                raise ValidationError("Offer Letter draft has not been generated for review.")

        # Safeguard: prevent accidental duplicate dispatch unless explicit
        if not force_resend:
            if loi_letter and loi_letter.status == LETTER_STATUS_SENT:
                raise ValidationError(
                    f"Letter of Intent has already been sent to candidate '{app.full_name}' ({app.email}) on "
                    f"{loi_letter.sent_at.strftime('%d %b %Y, %I:%M %p')}. To resend, please confirm explicit re-dispatch."
                )
            if offer_letter and offer_letter.status == LETTER_STATUS_SENT:
                raise ValidationError(
                    f"Offer Letter has already been sent to candidate '{app.full_name}' ({app.email}) on "
                    f"{offer_letter.sent_at.strftime('%d %b %Y, %I:%M %p')}. To resend, please confirm explicit re-dispatch."
                )

        # Validate mandatory placeholder fields: joining_date
        loi_joining = (loi_letter.field_values or {}).get("joining_date", "").strip() if loi_letter else ""
        offer_joining = (offer_letter.field_values or {}).get("joining_date", "").strip() if offer_letter else ""
        app_joining = (getattr(app, "joining_date", None) or "").strip()
        joining_str = offer_joining or loi_joining or app_joining

        # Automatically synchronize joining date so neither letter misses it
        if joining_str:
            if include_loi and loi_letter and not loi_joining:
                fv = dict(loi_letter.field_values or {})
                fv["joining_date"] = joining_str
                loi_letter.field_values = fv
                loi_joining = joining_str
            if include_offer and offer_letter and not offer_joining:
                fv = dict(offer_letter.field_values or {})
                fv["joining_date"] = joining_str
                offer_letter.field_values = fv
                offer_joining = joining_str

        if (include_loi and not loi_joining) or (include_offer and not offer_joining):
            raise ValidationError(
                "Missing mandatory field: 'Joining Date'. Please specify the candidate's joining date "
                "in the letter review screen before sending."
            )

        # Prepare attachments: generate and strictly validate candidate-specific PDFs
        attachments = []
        safe_cand_name = (app.full_name or "Candidate").replace(" ", "_")

        if include_loi and loi_letter:
            LetterGenerationService.ensure_candidate_docx_exists(loi_letter)
            loi_file = LetterGenerationService.get_or_render_pdf(loi_letter)
            if not loi_file or not loi_file.endswith(".pdf") or not os.path.exists(loi_file):
                raise ValidationError("Could not generate final PDF for Letter of Intent. Email dispatch aborted.")
            attachments.append({
                "filename": f"Letter_of_Intent_{safe_cand_name}.pdf",
                "filepath": loi_file,
            })

        if include_offer and offer_letter:
            LetterGenerationService.ensure_candidate_docx_exists(offer_letter)
            offer_file = LetterGenerationService.get_or_render_pdf(offer_letter)
            if not offer_file or not offer_file.endswith(".pdf") or not os.path.exists(offer_file):
                raise ValidationError("Could not generate final PDF for Offer Letter. Email dispatch aborted.")
            attachments.append({
                "filename": f"Offer_Letter_{safe_cand_name}.pdf",
                "filepath": offer_file,
            })

        if not attachments:
            raise ValidationError("No document attachments were selected for dispatch.")

        subject, html_body, text_body = EmailService.render_congratulations_email(
            candidate_name=app.full_name,
            application_number=app.application_number,
            registered_email=app.email,
            role_applied=app.role_applied or "Candidate",
            attachments=attachments,
        )

        try:
            send_result = EmailService.send_email(
                to_email=app.email,
                subject=subject,
                html_body=html_body,
                text_body=text_body,
                attachments=attachments,
            )

            now = utc_now()
            with atomic():
                if include_loi and loi_letter:
                    loi_letter.status = LETTER_STATUS_SENT
                    loi_letter.sent_at = now
                    loi_letter.sent_by_id = user.id
                    loi_letter.email_recipient = app.email
                    loi_letter.email_status = "SENT"
                    loi_letter.email_error = None

                if include_offer and offer_letter:
                    offer_letter.status = LETTER_STATUS_SENT
                    offer_letter.sent_at = now
                    offer_letter.sent_by_id = user.id
                    offer_letter.email_recipient = app.email
                    offer_letter.email_status = "SENT"
                    offer_letter.email_error = None

                audit = AuditLog(
                    action="LETTERS_SENT",
                    user_id=user.id,
                    resource="candidate_applications",
                    resource_id=str(app.id),
                    details={
                        "application_number": app.application_number,
                        "recipient_email": app.email,
                        "joining_date": joining_str,
                        "document_scope": document_scope,
                        "mode": send_result.get("mode"),
                        "message_id": send_result.get("message_id"),
                    },
                )
                db.session.add(audit)

            mode = send_result.get("mode")
            if mode == "smtp":
                msg_txt = f"Congratulation email and finalized document(s) sent successfully to '{app.email}' via live SMTP."
            else:
                msg_txt = (
                    f"Email packaged and saved to local storage in emulated mode ({send_result.get('log_path', '')}). "
                    "To send live emails, please configure SMTP_HOST, SMTP_PORT, SMTP_USERNAME, and SMTP_PASSWORD in backend/.env."
                )

            return {
                "success": True,
                "message": msg_txt,
                "recipient": app.email,
                "sent_at": now.isoformat(),
                "delivery_mode": mode,
                "document_scope": document_scope,
            }
        except Exception as e:
            with atomic():
                if include_loi and loi_letter:
                    loi_letter.email_status = "FAILED"
                    loi_letter.email_error = str(e)
                if include_offer and offer_letter:
                    offer_letter.email_status = "FAILED"
                    offer_letter.email_error = str(e)

            raise ValidationError(f"Failed to email letters to '{app.email}': {str(e)}")

    @staticmethod
    def preview_congratulations_email(
        application_id: uuid.UUID,
        user: Any,
        document_scope: str = "both",
    ) -> Dict[str, Any]:
        """Generate email preview payload for HR review prior to sending."""
        app = db.session.execute(
            select(CandidateApplication).where(CandidateApplication.id == application_id)
        ).scalar_one_or_none()

        if not app:
            raise NotFoundError("Application record not found.")

        if app.status != APPLICATION_STATUS_APPROVED:
            raise ValidationError(
                f"Candidate application '{app.application_number}' has status '{app.status}'. "
                "Letters can only be previewed for APPROVED candidates."
            )

        include_loi = document_scope in ("both", TEMPLATE_TYPE_LOI, "loi")
        include_offer = document_scope in ("both", TEMPLATE_TYPE_OFFER, "offer")

        attachments = []
        safe_cand_name = (app.full_name or "Candidate").replace(" ", "_")

        if include_loi:
            attachments.append({
                "filename": f"Letter_of_Intent_{safe_cand_name}.pdf",
            })
        if include_offer:
            attachments.append({
                "filename": f"Offer_Letter_{safe_cand_name}.pdf",
            })

        subject, html_body, text_body = EmailService.render_congratulations_email(
            candidate_name=app.full_name or "Candidate",
            application_number=app.application_number,
            registered_email=app.email or "",
            role_applied=app.role_applied or "Candidate",
            attachments=attachments,
        )

        return {
            "subject": subject,
            "html": html_body,
            "text": text_body,
            "recipient": app.email,
            "candidate_name": app.full_name,
            "application_number": app.application_number,
            "role_applied": app.role_applied or "Candidate",
            "attachments": [a["filename"] for a in attachments],
        }

    @staticmethod
    def accept_candidate_offer(
        application_id: uuid.UUID,
        user: User,
        ip_address: Optional[str] = None,
        user_agent: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Candidate accepts the formal offer letter and agreement.
        
        Saves acceptance timestamp and metadata, validates authorization,
        prevents duplicate submissions, and logs the event to AuditLog.
        """
        app = db.session.execute(
            select(CandidateApplication).where(CandidateApplication.id == application_id)
        ).scalar_one_or_none()

        if not app:
            raise NotFoundError("Application record not found.")

        # Security check: User must be the candidate who owns the application or an authorized admin
        user_roles = getattr(user, "role_names", []) or []
        is_candidate_owner = (app.user_id == user.id)
        if not is_candidate_owner and "admin" not in user_roles:
            raise ForbiddenError("You are not authorized to accept this offer.")

        if app.status not in (APPLICATION_STATUS_APPROVED, APPLICATION_STATUS_OFFER_ACCEPTED):
            raise ValidationError(
                f"Candidate application '{app.application_number}' has status '{app.status}'. "
                "Offers can only be accepted for approved applications."
            )

        # Retrieve offer letter
        offer_letter = db.session.execute(
            select(CandidateLetter)
            .where(CandidateLetter.application_id == app.id)
            .where(CandidateLetter.letter_type == TEMPLATE_TYPE_OFFER)
            .order_by(CandidateLetter.version.desc(), CandidateLetter.created_at.desc())
        ).scalars().first()

        if not offer_letter:
            raise ValidationError("No formal Offer Letter found awaiting acceptance for this candidate.")

        # Check if already accepted to prevent duplicate events and repeat writes
        if offer_letter.status == LETTER_STATUS_ACCEPTED or app.offer_accepted_at is not None:
            accepted_dt = offer_letter.accepted_at or app.offer_accepted_at
            accepted_time_str = accepted_dt.strftime("%d %b %Y, %I:%M %p") if accepted_dt else "previously"
            return {
                "success": True,
                "already_accepted": True,
                "message": f"This offer has already been accepted on {accepted_time_str}.",
                "application_id": str(app.id),
                "application_number": app.application_number,
                "candidate_name": app.full_name,
                "offer_reference": (offer_letter.field_values or {}).get("reference_number") or app.application_number,
                "status": "ACCEPTED",
                "accepted_at": accepted_dt.isoformat() if accepted_dt else utc_now().isoformat(),
            }

        now = utc_now()
        offer_ref = (offer_letter.field_values or {}).get("reference_number") or app.application_number

        with atomic():
            # Update offer letter
            offer_letter.status = LETTER_STATUS_ACCEPTED
            offer_letter.accepted_at = now
            offer_letter.accepted_by_id = user.id
            offer_letter.accepted_ip = ip_address
            offer_letter.accepted_user_agent = user_agent

            # Also update LOI status if present
            loi_letter = db.session.execute(
                select(CandidateLetter)
                .where(CandidateLetter.application_id == app.id)
                .where(CandidateLetter.letter_type == TEMPLATE_TYPE_LOI)
                .order_by(CandidateLetter.version.desc(), CandidateLetter.created_at.desc())
            ).scalars().first()
            if loi_letter and loi_letter.status != LETTER_STATUS_ACCEPTED:
                loi_letter.status = LETTER_STATUS_ACCEPTED
                loi_letter.accepted_at = now
                loi_letter.accepted_by_id = user.id
                loi_letter.accepted_ip = ip_address
                loi_letter.accepted_user_agent = user_agent

            # Update application record
            app.offer_accepted_at = now
            app.offer_accepted_by_id = user.id
            app.offer_acceptance_status = "ACCEPTED"
            app.status = APPLICATION_STATUS_OFFER_ACCEPTED

            # Write server-side audit event matching requirements
            audit = AuditLog(
                action="CANDIDATE_OFFER_ACCEPTED",
                user_id=user.id,
                resource="candidate_applications",
                resource_id=str(app.id),
                ip_address=ip_address,
                user_agent=user_agent,
                details={
                    "activity": "Candidate Accepted Offer",
                    "status": "Accepted",
                    "candidate_name": app.full_name,
                    "application_number": app.application_number,
                    "offer_reference": offer_ref,
                    "letter_id": str(offer_letter.id),
                    "accepted_at": now.isoformat(),
                    "actor_email": user.email,
                    "actor_name": user.full_name or app.full_name,
                },
            )
            db.session.add(audit)

        return {
            "success": True,
            "already_accepted": False,
            "message": "Offer accepted successfully. Welcome to the team!",
            "application_id": str(app.id),
            "application_number": app.application_number,
            "candidate_name": app.full_name,
            "offer_reference": offer_ref,
            "status": "ACCEPTED",
            "accepted_at": now.isoformat(),
        }


